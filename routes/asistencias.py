# routes/asistencias.py
from __future__ import annotations

import calendar
import logging
from datetime import date, datetime, timedelta
from io import BytesIO

import openpyxl
from fastapi import APIRouter, Depends, Request
from fastapi.responses import JSONResponse, RedirectResponse, Response
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import and_, case, func, or_

from config import Config
from dependencies import get_current_user_id
from models import Aula, Asistencia, ConfiguracionPension, DetalleAsistenciaAula, Estudiante, Matricula, PensionEstudiante, db
from services.asistencia_service import AsistenciaService
from template_helpers import add_flash, common_context, csrf_ok, templates, url_for as app_url_for
from utils.datetime_helper import now_lima
from utils.email_helper import enviar_notificacion_asistencia_async
from utils.helpers import obtener_meses_pendientes
from utils.pagination import paginate_query

router = APIRouter()
logger = logging.getLogger(__name__)


def get_asistencia_service():
    return AsistenciaService


def _username(request: Request) -> str | None:
    return request.session.get("username")


def _static_photo_url(request: Request, estudiante: Estudiante) -> str | None:
    if not estudiante.foto_perfil:
        return None
    if estudiante.foto_perfil.startswith("uploads/"):
        return app_url_for(request, "static", filename=estudiante.foto_perfil)
    return app_url_for(request, "static", filename=f"uploads/estudiantes/{estudiante.foto_perfil}")


def _obtener_info_pension(estudiante_id):
    """Obtiene el estado de pensión de un estudiante para el año escolar actual."""
    try:
        config = ConfiguracionPension.query.filter_by(activo=True).first()
        anio = config.anio_escolar if config else str(now_lima().year)

        pension = PensionEstudiante.query.filter_by(
            estudiante_id=estudiante_id, anio_escolar=anio, activo=True
        ).first()

        if not pension:
            return {"asignada": False}

        orden_meses = [
            "enero",
            "febrero",
            "marzo",
            "abril",
            "mayo",
            "junio",
            "julio",
            "agosto",
            "septiembre",
            "octubre",
            "noviembre",
            "diciembre",
        ]
        mes_actual_idx = now_lima().month - 1

        todos_pendientes = obtener_meses_pendientes(estudiante_id, anio)
        meses = [
            m
            for m in todos_pendientes
            if m.lower() in orden_meses and orden_meses.index(m.lower()) <= mes_actual_idx
        ]

        return {
            "asignada": True,
            "tiene_deuda": len(meses) > 0,
            "meses_pendientes": meses,
            "cantidad": len(meses),
            "monto_mensual": float(pension.monto_mensual),
        }
    except Exception:
        return {"asignada": False}


@router.get("/scan", name="asistencias.scan")
def scan(request: Request, _user_id: int = Depends(get_current_user_id)):
    return templates.TemplateResponse("asistencias/scan.html", common_context(request))


@router.post("/api/scan", name="asistencias.api_scan")
async def api_scan(request: Request, _user_id: int = Depends(get_current_user_id)):
    try:
        try:
            payload = await request.json()
        except Exception:
            payload = {}
        if not isinstance(payload, dict):
            payload = {}
        code = (payload.get("code") or "").strip()

        if not code:
            return JSONResponse({"ok": False, "message": "Código vacío."}, status_code=400)

        estudiante = Estudiante.query.filter_by(codigo_estudiante=code).first()
        if not estudiante:
            estudiante = Estudiante.query.filter_by(dni_est=code).first()

        if not estudiante:
            return JSONResponse(
                {"ok": False, "message": f"No existe estudiante con código {code}."},
                status_code=404,
            )

        now = now_lima()
        hora_actual = now.hour
        tipo = "ENTRADA" if hora_actual < 13 else "SALIDA"

        hoy = date.today()
        asistencia_existente = (
            Asistencia.query.filter(
                and_(
                    Asistencia.estudiante_id == estudiante.id,
                    Asistencia.tipo == tipo,
                    func.date(Asistencia.fecha_hora) == hoy,
                )
            ).first()
        )

        if asistencia_existente:
            tipo_texto = "entrada" if tipo == "ENTRADA" else "salida"
            ts = now.strftime("%Y-%m-%d %H:%M:%S")
            try:
                from utils.email_helper import enviar_notificacion_duplicado_async

                enviar_notificacion_duplicado_async(
                    estudiante,
                    tipo,
                    now.strftime("%H:%M:%S"),
                    asistencia_existente.fecha_hora.strftime("%H:%M:%S"),
                )
            except Exception as e:
                logger.error("Error al enviar notificación de duplicado: %s", e, exc_info=True)

            photo_url = _static_photo_url(request, estudiante)
            return JSONResponse(
                {
                    "ok": False,
                    "message": (
                        f"{estudiante.nombres_est} {estudiante.apellido_paterno_est} "
                        f"ya registró su {tipo_texto} el día de hoy."
                    ),
                    "when": ts,
                    "tipo": tipo,
                    "student": {
                        "nombres": estudiante.nombres_est,
                        "apellidos": f"{estudiante.apellido_paterno_est} {estudiante.apellido_materno_est}",
                        "codigo": estudiante.codigo_estudiante or "",
                        "photo_url": photo_url,
                    },
                    "pension": _obtener_info_pension(estudiante.id),
                },
                status_code=400,
            )

        nueva_asistencia = Asistencia(estudiante_id=estudiante.id, tipo=tipo, fecha_hora=now)
        db.session.add(nueva_asistencia)
        db.session.commit()

        ts = now.strftime("%Y-%m-%d %H:%M:%S")
        try:
            enviar_notificacion_asistencia_async(estudiante, tipo, now.strftime("%H:%M:%S"))
        except Exception as e:
            logger.error("Error al programar envío de email: %s", e, exc_info=True)

        photo_url = _static_photo_url(request, estudiante)
        tipo_texto = "Entrada" if tipo == "ENTRADA" else "Salida"
        return JSONResponse(
            {
                "ok": True,
                "message": (
                    f"{tipo_texto} registrada: {estudiante.nombres_est} "
                    f"{estudiante.apellido_paterno_est} - {ts}"
                ),
                "attendance_id": nueva_asistencia.id,
                "when": ts,
                "tipo": tipo,
                "student": {
                    "nombres": estudiante.nombres_est,
                    "apellidos": f"{estudiante.apellido_paterno_est} {estudiante.apellido_materno_est}",
                    "codigo": estudiante.codigo_estudiante or "",
                    "photo_url": photo_url,
                },
                "pension": _obtener_info_pension(estudiante.id),
            }
        )

    except Exception as e:
        db.session.rollback()
        print(f"Error al registrar asistencia: {e}")
        return JSONResponse({"ok": False, "message": "Error al registrar asistencia"}, status_code=500)


@router.get("/dashboard", name="asistencias.dashboard")
def dashboard(request: Request, _user_id: int = Depends(get_current_user_id)):
    hoy = date.today()
    hace_7_dias = hoy - timedelta(days=6)

    def _empty_stats():
        return {
            "total_estudiantes": 0,
            "total_asistencias": 0,
            "asistencias_hoy": 0,
            "asistencias_entrada": 0,
            "asistencias_salida": 0,
            "asistencias_por_dia": [
                {"fecha": (hoy - timedelta(days=i)).strftime("%Y-%m-%d"), "count": 0}
                for i in range(6, -1, -1)
            ],
            "top_estudiantes": [],
        }

    try:
        q = (request.query_params.get("q") or "").strip()

        query = db.session.query(Asistencia, Estudiante).join(
            Estudiante, Asistencia.estudiante_id == Estudiante.id
        )

        if q:
            query = query.filter(
                or_(
                    Estudiante.nombres_est.ilike(f"%{q}%"),
                    Estudiante.apellido_paterno_est.ilike(f"%{q}%"),
                    Estudiante.apellido_materno_est.ilike(f"%{q}%"),
                    Estudiante.dni_est.ilike(f"%{q}%"),
                )
            )

        page = int(request.query_params.get("page") or 1)
        ordered = query.order_by(Asistencia.fecha_hora.desc())
        paginacion = paginate_query(ordered, page=page, per_page=50, error_out=False)
        asistencias = paginacion.items

        # Usamos rangos de datetime (portable Postgres/SQLite) en lugar de func.date() == str(hoy),
        # que puede fallar en Postgres por mismatch date/varchar.
        inicio_hoy = datetime.combine(hoy, datetime.min.time())
        inicio_manana = datetime.combine(hoy + timedelta(days=1), datetime.min.time())

        stats_row = db.session.query(
            func.count(Asistencia.id),
            func.sum(
                case(
                    (
                        and_(
                            Asistencia.fecha_hora >= inicio_hoy,
                            Asistencia.fecha_hora < inicio_manana,
                        ),
                        1,
                    ),
                    else_=0,
                )
            ),
            func.sum(case((Asistencia.tipo == "ENTRADA", 1), else_=0)),
            func.sum(case((Asistencia.tipo == "SALIDA", 1), else_=0)),
        ).first()

        total_estudiantes = Estudiante.query.count()

        asistencias_agrupadas = (
            db.session.query(
                func.date(Asistencia.fecha_hora).label("fecha"),
                func.count(Asistencia.id).label("total"),
            )
            .filter(Asistencia.fecha_hora >= datetime.combine(hace_7_dias, datetime.min.time()))
            .group_by(func.date(Asistencia.fecha_hora))
            .all()
        )

        dias_dict = {}
        for row in asistencias_agrupadas:
            fecha_val = row.fecha
            if hasattr(fecha_val, "strftime"):
                key = fecha_val.strftime("%Y-%m-%d")
            else:
                key = str(fecha_val)
            dias_dict[key] = row.total

        asistencias_por_dia = []
        for i in range(6, -1, -1):
            fecha = hoy - timedelta(days=i)
            asistencias_por_dia.append(
                {"fecha": fecha.strftime("%Y-%m-%d"), "count": dias_dict.get(fecha.strftime("%Y-%m-%d"), 0)}
            )

        top_estudiantes = (
            db.session.query(Estudiante, func.count(Asistencia.id).label("total"))
            .join(Asistencia, Asistencia.estudiante_id == Estudiante.id)
            .group_by(Estudiante.id)
            .order_by(func.count(Asistencia.id).desc())
            .limit(5)
            .all()
        )

        stats = {
            "total_estudiantes": total_estudiantes,
            "total_asistencias": stats_row[0] or 0,
            "asistencias_hoy": stats_row[1] or 0,
            "asistencias_entrada": stats_row[2] or 0,
            "asistencias_salida": stats_row[3] or 0,
            "asistencias_por_dia": asistencias_por_dia,
            "top_estudiantes": [
                {
                    "nombre": (
                        f"{estudiante.nombres_est} {estudiante.apellido_paterno_est} "
                        f"{estudiante.apellido_materno_est}"
                    ),
                    "total": total,
                }
                for estudiante, total in top_estudiantes
            ],
        }

        return templates.TemplateResponse(
            "asistencias/dashboard.html",
            common_context(request, asistencias=asistencias, paginacion=paginacion, stats=stats, q=q),
        )

    except Exception as e:
        logger.exception("Error al cargar dashboard de asistencias: %s", e)
        try:
            db.session.rollback()
        except Exception:
            pass
        add_flash(request, "Error al cargar el dashboard de asistencias", "error")
        return templates.TemplateResponse(
            "asistencias/dashboard.html",
            common_context(
                request,
                asistencias=[],
                paginacion=None,
                stats=_empty_stats(),
                q=(request.query_params.get("q") or "").strip(),
            ),
        )


@router.api_route("/reportes", methods=["GET", "POST"], name="asistencias.reportes")
async def reportes(request: Request, _user_id: int = Depends(get_current_user_id)):
    if request.method == "POST":
        form = await request.form()
        if getattr(Config, "WTF_CSRF_ENABLED", True) and not csrf_ok(request, form.get("csrf_token")):
            add_flash(request, "Sesión de seguridad expirada. Intente de nuevo.", "error")
            return RedirectResponse(url=str(request.url_for("asistencias.reportes")), status_code=303)
        try:
            fecha_desde = (form.get("fecha_desde") or "").strip()
            fecha_hasta = (form.get("fecha_hasta") or "").strip()
            tipo = (form.get("tipo") or "").strip()
            nivel = (form.get("nivel") or "").strip()
            grado = (form.get("grado") or "").strip()
            seccion = (form.get("seccion") or "").strip()

            query = db.session.query(Asistencia, Estudiante).join(
                Estudiante, Asistencia.estudiante_id == Estudiante.id
            )

            if fecha_desde:
                query = query.filter(func.date(Asistencia.fecha_hora) >= fecha_desde)
            if fecha_hasta:
                query = query.filter(func.date(Asistencia.fecha_hora) <= fecha_hasta)
            if tipo:
                query = query.filter(Asistencia.tipo == tipo)

            # Filtrado por nivel/grado/seccion: lo resolvemos vía matrícula
            # activa en aulas que coincidan (no por Estudiante.nivel/grado,
            # que pueden estar desactualizados respecto al aula real).
            if nivel or grado or seccion:
                aula_q = Aula.query.filter(Aula.activo.is_(True))
                if nivel:
                    aula_q = aula_q.filter(func.upper(Aula.nivel) == nivel.upper())
                if grado:
                    aula_q = aula_q.filter(Aula.grado == grado)
                if seccion:
                    aula_q = aula_q.filter(Aula.seccion == seccion)
                aula_ids_match = [a.id for a in aula_q.all()]
                if aula_ids_match:
                    est_ids_match = [
                        r[0]
                        for r in (
                            db.session.query(Matricula.estudiante_id)
                            .filter(
                                Matricula.aula_id.in_(aula_ids_match),
                                Matricula.estado == "activo",
                            )
                            .distinct()
                            .all()
                        )
                    ]
                    if est_ids_match:
                        query = query.filter(Estudiante.id.in_(est_ids_match))
                    else:
                        query = query.filter(False)
                else:
                    query = query.filter(False)

            asistencias_rows = query.order_by(Asistencia.fecha_hora.desc()).limit(10000).all()

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Reporte de Asistencia"

            headers = ["ID", "Nombres", "Apellidos", "DNI", "Nivel", "Grado", "Sección", "Tipo", "Fecha", "Hora"]
            ws.append(headers)

            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")

            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Precargar matrícula activa (nivel/grado/sección real) por estudiante
            est_ids = list({e.id for _, e in asistencias_rows})
            mats_por_est: dict = {}
            if est_ids:
                for m in (
                    db.session.query(Matricula, Aula)
                    .join(Aula, Matricula.aula_id == Aula.id)
                    .filter(
                        Matricula.estudiante_id.in_(est_ids),
                        Matricula.estado == "activo",
                    )
                    .all()
                ):
                    mat, aula = m
                    mats_por_est[mat.estudiante_id] = aula

            for asistencia, estudiante in asistencias_rows:
                aula_mat = mats_por_est.get(estudiante.id)
                if aula_mat is not None:
                    nivel_txt = aula_mat.nivel
                    grado_txt = aula_mat.grado
                    sec_txt = aula_mat.seccion or "-"
                else:
                    nivel_txt = estudiante.nivel or "-"
                    grado_txt = estudiante.grado or "-"
                    sec_txt = estudiante.seccion or "-"
                ws.append(
                    [
                        asistencia.id,
                        estudiante.nombres_est,
                        f"{estudiante.apellido_paterno_est} {estudiante.apellido_materno_est}",
                        estudiante.dni_est,
                        nivel_txt,
                        grado_txt,
                        sec_txt,
                        asistencia.tipo,
                        asistencia.fecha_str,
                        asistencia.hora_str,
                    ]
                )

            for col, w in zip(
                ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J"],
                [8, 20, 25, 12, 14, 14, 12, 12, 12, 12],
            ):
                ws.column_dimensions[col].width = w

            output = BytesIO()
            wb.save(output)
            output.seek(0)

            nombre_base = "reporte_asistencia"
            if nivel:
                nombre_base += f"_{nivel}"
            if grado:
                nombre_base += f"_{grado}"
            if seccion:
                nombre_base += f"_{seccion}"
            filename = f"{nombre_base}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

            return Response(
                output.getvalue(),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": f"attachment; filename={filename}"},
            )

        except Exception as e:
            print(f"Error al generar reporte: {e}")
            add_flash(request, f"Error al generar el reporte: {str(e)}", "error")
            return RedirectResponse(url=str(request.url_for("asistencias.reportes")), status_code=303)

    total_asistencias = Asistencia.query.count()
    total_estudiantes = Estudiante.query.count()

    # ──────────────────────────────────────────────────────────────────
    # Datos de filtros dinámicos: sólo desde aulas activas reales.
    # Sigue el mismo patrón que /asistencias/por_aula y /form_estudiante
    # para que los niveles, grados y programas reflejen sólo las aulas
    # actualmente configuradas (no los campos históricos de Estudiante).
    # ──────────────────────────────────────────────────────────────────
    aulas_activas = (
        Aula.query.filter(Aula.activo.is_(True))
        .order_by(Aula.nivel, Aula.grado, Aula.seccion)
        .all()
    )

    niveles = sorted({a.nivel for a in aulas_activas if a.nivel})

    grados_por_nivel: dict[str, list[str]] = {}
    for a in aulas_activas:
        if not a.nivel or not a.grado:
            continue
        key = a.nivel.upper()
        bucket = grados_por_nivel.setdefault(key, [])
        if a.grado not in bucket:
            bucket.append(a.grado)

    secciones_por_nivel_grado: dict[str, list[str]] = {}
    for a in aulas_activas:
        if not a.nivel or not a.grado or not a.seccion:
            continue
        key = f"{a.nivel.upper()}|{a.grado}"
        bucket = secciones_por_nivel_grado.setdefault(key, [])
        if a.seccion not in bucket:
            bucket.append(a.seccion)

    return templates.TemplateResponse(
        "asistencias/reportes.html",
        common_context(
            request,
            total_asistencias=total_asistencias,
            total_estudiantes=total_estudiantes,
            niveles=niveles,
            grados_por_nivel=grados_por_nivel,
            secciones_por_nivel_grado=secciones_por_nivel_grado,
        ),
    )


@router.get("/asistencias/reporte-mensual", name="asistencias.reporte_mensual")
def reporte_mensual(request: Request, _user_id: int = Depends(get_current_user_id)):
    nivel = request.query_params.get("nivel", "")
    grado = request.query_params.get("grado", "")
    seccion = request.query_params.get("seccion", "")
    mes_raw = request.query_params.get("mes", str(datetime.now().month))
    anio_raw = request.query_params.get("anio", str(datetime.now().year))
    page = int(request.query_params.get("page") or 1)
    per_page = 20

    try:
        mes = int(mes_raw)
        anio = int(anio_raw)
    except ValueError:
        mes = datetime.now().month
        anio = datetime.now().year

    query = Estudiante.query
    if nivel:
        query = query.filter_by(nivel=nivel)
    if grado:
        query = query.filter_by(grado=grado)
    if seccion:
        query = query.filter_by(seccion=seccion)

    ordered = query.order_by(Estudiante.apellido_paterno_est, Estudiante.nombres_est)
    estudiantes = paginate_query(ordered, page=page, per_page=per_page, error_out=False)

    num_dias = calendar.monthrange(anio, mes)[1]
    dias_mes = list(range(1, num_dias + 1))

    fecha_inicio = date(anio, mes, 1)
    fecha_fin = date(anio, mes, num_dias)

    estudiante_ids = [e.id for e in estudiantes.items]

    if estudiante_ids:
        asistencias = (
            Asistencia.query.filter(
                and_(
                    Asistencia.estudiante_id.in_(estudiante_ids),
                    func.date(Asistencia.fecha_hora) >= fecha_inicio,
                    func.date(Asistencia.fecha_hora) <= fecha_fin,
                )
            )
            .all()
        )
    else:
        asistencias = []

    asistencias_dict = {}
    for asistencia in asistencias:
        key = (asistencia.estudiante_id, asistencia.fecha_hora.day)
        if key not in asistencias_dict:
            asistencias_dict[key] = []
        asistencias_dict[key].append(asistencia.tipo)

    niveles = [
        n[0]
        for n in db.session.query(Estudiante.nivel)
        .distinct()
        .filter(Estudiante.nivel.isnot(None))
        .order_by(Estudiante.nivel)
        .all()
    ]
    grados = [
        g[0]
        for g in db.session.query(Estudiante.grado)
        .distinct()
        .filter(Estudiante.grado.isnot(None))
        .order_by(Estudiante.grado)
        .all()
    ]
    secciones = [
        s[0]
        for s in db.session.query(Estudiante.seccion)
        .distinct()
        .filter(Estudiante.seccion.isnot(None))
        .order_by(Estudiante.seccion)
        .all()
    ]

    meses_nombres = [
        "",
        "Enero",
        "Febrero",
        "Marzo",
        "Abril",
        "Mayo",
        "Junio",
        "Julio",
        "Agosto",
        "Septiembre",
        "Octubre",
        "Noviembre",
        "Diciembre",
    ]
    mes_nombre = meses_nombres[mes]

    return templates.TemplateResponse(
        "asistencias/reporte_mensual.html",
        common_context(
            request,
            estudiantes=estudiantes,
            dias_mes=dias_mes,
            asistencias_dict=asistencias_dict,
            nivel=nivel,
            grado=grado,
            seccion=seccion,
            mes=mes,
            anio=anio,
            mes_nombre=mes_nombre,
            niveles=niveles,
            grados=grados,
            secciones=secciones,
        ),
    )


@router.get("/asistencias/reporte-mensual/excel", name="asistencias.reporte_mensual_excel")
def reporte_mensual_excel(request: Request, _user_id: int = Depends(get_current_user_id)):
    nivel = request.query_params.get("nivel", "")
    grado = request.query_params.get("grado", "")
    seccion = request.query_params.get("seccion", "")
    mes_raw = request.query_params.get("mes", str(datetime.now().month))
    anio_raw = request.query_params.get("anio", str(datetime.now().year))

    try:
        mes = int(mes_raw)
        anio = int(anio_raw)
    except ValueError:
        mes = datetime.now().month
        anio = datetime.now().year

    meses_nombres = [
        "",
        "Enero",
        "Febrero",
        "Marzo",
        "Abril",
        "Mayo",
        "Junio",
        "Julio",
        "Agosto",
        "Septiembre",
        "Octubre",
        "Noviembre",
        "Diciembre",
    ]
    mes_nombre = meses_nombres[mes]

    query = Estudiante.query
    if nivel:
        query = query.filter_by(nivel=nivel)
    if grado:
        query = query.filter_by(grado=grado)
    if seccion:
        query = query.filter_by(seccion=seccion)

    estudiantes = query.order_by(Estudiante.apellido_paterno_est, Estudiante.nombres_est).all()

    num_dias = calendar.monthrange(anio, mes)[1]
    dias_mes = list(range(1, num_dias + 1))

    fecha_inicio = date(anio, mes, 1)
    fecha_fin = date(anio, mes, num_dias)

    estudiante_ids = [e.id for e in estudiantes]

    if estudiante_ids:
        asistencias = (
            Asistencia.query.filter(
                and_(
                    Asistencia.estudiante_id.in_(estudiante_ids),
                    func.date(Asistencia.fecha_hora) >= fecha_inicio,
                    func.date(Asistencia.fecha_hora) <= fecha_fin,
                )
            )
            .all()
        )
    else:
        asistencias = []

    asistencias_dict = {}
    for asistencia in asistencias:
        key = (asistencia.estudiante_id, asistencia.fecha_hora.day)
        if key not in asistencias_dict:
            asistencias_dict[key] = []
        asistencias_dict[key].append(asistencia.tipo)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Asistencias {mes_nombre} {anio}"

    filtro_texto = f"{mes_nombre} {anio}"
    if nivel:
        filtro_texto += f" - {nivel}"
    if grado:
        filtro_texto += f" {grado}"
    if seccion:
        filtro_texto += f" {seccion}"

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5 + len(dias_mes))
    title_cell = ws.cell(row=1, column=1, value=f"Reporte Mensual de Asistencias - {filtro_texto}")
    title_cell.font = Font(bold=True, size=14, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="5F2A5D", end_color="5F2A5D", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    headers = ["N°", "Apellidos", "Nombres", "DNI", "Nivel", "Grado"]
    for dia in dias_mes:
        headers.append(str(dia))
    headers.extend(["Asistencias", "Faltas"])

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill(start_color="7c3b7a", end_color="7c3b7a", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    fill_completo = PatternFill(start_color="22C55E", end_color="22C55E", fill_type="solid")
    fill_entrada = PatternFill(start_color="EAB308", end_color="EAB308", fill_type="solid")
    fill_salida = PatternFill(start_color="F97316", end_color="F97316", fill_type="solid")
    fill_falta = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    font_white = Font(color="FFFFFF", size=9, bold=True)
    font_red = Font(color="DC2626", size=9, bold=True)
    center_align = Alignment(horizontal="center", vertical="center")

    for idx, estudiante in enumerate(estudiantes, 1):
        row = idx + 3
        ws.cell(row=row, column=1, value=idx).alignment = center_align
        ws.cell(row=row, column=2, value=f"{estudiante.apellido_paterno_est} {estudiante.apellido_materno_est}")
        ws.cell(row=row, column=3, value=estudiante.nombres_est)
        ws.cell(row=row, column=4, value=estudiante.dni_est or "-").alignment = center_align
        ws.cell(row=row, column=5, value=estudiante.nivel or "-").alignment = center_align
        ws.cell(row=row, column=6, value=estudiante.grado or "-").alignment = center_align

        total_asistencias = 0.0
        total_faltas = 0.0

        for dia_idx, dia in enumerate(dias_mes):
            col = 7 + dia_idx
            asistencias_dia = asistencias_dict.get((estudiante.id, dia), [])
            cell = ws.cell(row=row, column=col)
            cell.alignment = center_align

            if asistencias_dia:
                if "ENTRADA" in asistencias_dia and "SALIDA" in asistencias_dia:
                    cell.value = "A"
                    cell.fill = fill_completo
                    cell.font = font_white
                    total_asistencias += 1
                elif "ENTRADA" in asistencias_dia:
                    cell.value = "E"
                    cell.fill = fill_entrada
                    cell.font = font_white
                    total_asistencias += 0.5
                    total_faltas += 0.5
                elif "SALIDA" in asistencias_dia:
                    cell.value = "S"
                    cell.fill = fill_salida
                    cell.font = font_white
                    total_asistencias += 0.5
                    total_faltas += 0.5
            else:
                cell.value = "F"
                cell.fill = fill_falta
                cell.font = font_red
                total_faltas += 1

        col_asist = 7 + len(dias_mes)
        col_faltas = col_asist + 1
        cell_asist = ws.cell(row=row, column=col_asist, value=total_asistencias)
        cell_asist.alignment = center_align
        cell_asist.font = Font(bold=True, color="16A34A")

        cell_faltas = ws.cell(row=row, column=col_faltas, value=total_faltas)
        cell_faltas.alignment = center_align
        cell_faltas.font = Font(bold=True, color="DC2626")

    row_leyenda = len(estudiantes) + 5
    ws.cell(row=row_leyenda, column=1, value="Leyenda:").font = Font(bold=True, size=10)
    leyendas = [
        ("A = Asistió (Entrada y Salida)", fill_completo),
        ("E = Solo Entrada", fill_entrada),
        ("S = Solo Salida", fill_salida),
        ("F = Faltó", fill_falta),
    ]
    for i, (texto, fill) in enumerate(leyendas):
        cell_color = ws.cell(row=row_leyenda + 1 + i, column=1)
        cell_color.fill = fill
        cell_color.value = "  "
        ws.cell(row=row_leyenda + 1 + i, column=2, value=texto)

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 10
    for dia_idx in range(len(dias_mes)):
        col_letter = openpyxl.utils.get_column_letter(7 + dia_idx)
        ws.column_dimensions[col_letter].width = 4
    col_asist_letter = openpyxl.utils.get_column_letter(7 + len(dias_mes))
    col_faltas_letter = openpyxl.utils.get_column_letter(8 + len(dias_mes))
    ws.column_dimensions[col_asist_letter].width = 12
    ws.column_dimensions[col_faltas_letter].width = 10

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    nombre_base = f"reporte_mensual_{mes_nombre}_{anio}"
    if nivel:
        nombre_base += f"_{nivel}"
    if grado:
        nombre_base += f"_{grado}"
    if seccion:
        nombre_base += f"_{seccion}"
    filename = f"{nombre_base}.xlsx"

    return Response(
        output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/por_aula", name="asistencias.por_aula")
def por_aula(request: Request, _user_id: int = Depends(get_current_user_id)):
    fecha_inicio = request.query_params.get("fecha_inicio", "")
    fecha_fin = request.query_params.get("fecha_fin", "")
    aula_id = request.query_params.get("aula_id", "")
    estado = request.query_params.get("estado", "")
    nivel = request.query_params.get("nivel", "")
    grado = request.query_params.get("grado", "")

    fecha_inicio_obj = (
        datetime.strptime(fecha_inicio, "%Y-%m-%d").date() if fecha_inicio else None
    )
    fecha_fin_obj = datetime.strptime(fecha_fin, "%Y-%m-%d").date() if fecha_fin else None

    todas_aulas = Aula.query.filter_by(activo=True).order_by(Aula.nivel, Aula.nombre).all()
    aulas_por_nivel = {}
    for aula in todas_aulas:
        aulas_por_nivel.setdefault(aula.nivel, []).append(aula)

    # Programas de Academia: únicamente los que existen como aulas activas.
    # (Ya no se incluyen grados heredados de estudiantes previamente matriculados,
    # para que el filtro refleje solo la configuración actual de aulas.)
    grados_academia = sorted(
        {a.grado for a in todas_aulas if a.nivel.upper() == "ACADEMIA" and a.grado}
    )

    nivel_u = (nivel or "").upper()
    es_academia_grado = nivel_u == "ACADEMIA" and bool(grado)
    registros_academia = []
    estudiantes_academia: list[dict] = []
    total_dias_academia = 0
    registros = []

    if es_academia_grado:
        # Buscar estudiantes vía matrícula activa en aulas de Academia
        # con el programa/grado seleccionado (no por Estudiante.grado,
        # que puede estar desactualizado respecto al aula donde está
        # realmente matriculado).
        aulas_academia_match = (
            Aula.query.filter(
                Aula.activo.is_(True),
                func.upper(Aula.nivel) == "ACADEMIA",
                Aula.grado == grado,
            )
            .all()
        )
        aula_ids_academia = [a.id for a in aulas_academia_match]

        if aula_ids_academia:
            matriculas_academia = (
                Matricula.query.filter(
                    Matricula.aula_id.in_(aula_ids_academia),
                    Matricula.estado == "activo",
                )
                .all()
            )
            ids_est_mat = [m.estudiante_id for m in matriculas_academia]
            est_grado = (
                Estudiante.query.filter(Estudiante.id.in_(ids_est_mat))
                .order_by(
                    Estudiante.apellido_paterno_est,
                    Estudiante.apellido_materno_est,
                    Estudiante.nombres_est,
                )
                .all()
            )
        else:
            est_grado = []
        ids_est = [e.id for e in est_grado]
        total_est = len(ids_est)

        if ids_est:
            q = (
                db.session.query(
                    func.date(Asistencia.fecha_hora).label("fecha"),
                    func.count(func.distinct(Asistencia.estudiante_id)).label("presentes"),
                )
                .filter(
                    Asistencia.estudiante_id.in_(ids_est),
                    Asistencia.tipo == "ENTRADA",
                )
                .group_by(func.date(Asistencia.fecha_hora))
            )

            if fecha_inicio_obj:
                q = q.filter(func.date(Asistencia.fecha_hora) >= fecha_inicio_obj)
            if fecha_fin_obj:
                q = q.filter(func.date(Asistencia.fecha_hora) <= fecha_fin_obj)

            fechas_clase: set[date] = set()
            for row in q.order_by(func.date(Asistencia.fecha_hora).desc()).all():
                fecha_val = row.fecha
                if not isinstance(fecha_val, date):
                    fecha_val = date.fromisoformat(str(fecha_val))
                registros_academia.append(
                    {
                        "fecha": fecha_val,
                        "grado": grado,
                        "total_estudiantes": total_est,
                        "total_presentes": row.presentes,
                        "total_ausentes": max(total_est - row.presentes, 0),
                    }
                )
                fechas_clase.add(fecha_val)
            total_dias_academia = len(fechas_clase)

            q_est = db.session.query(
                Asistencia.estudiante_id,
                func.count(func.distinct(func.date(Asistencia.fecha_hora))).label("dias_presentes"),
                func.max(Asistencia.fecha_hora).label("ultima_asistencia"),
            ).filter(
                Asistencia.estudiante_id.in_(ids_est),
                Asistencia.tipo == "ENTRADA",
            )
            if fecha_inicio_obj:
                q_est = q_est.filter(func.date(Asistencia.fecha_hora) >= fecha_inicio_obj)
            if fecha_fin_obj:
                q_est = q_est.filter(func.date(Asistencia.fecha_hora) <= fecha_fin_obj)

            stats_por_est = {
                eid: {"dias_presentes": dp or 0, "ultima": ult}
                for eid, dp, ult in q_est.group_by(Asistencia.estudiante_id).all()
            }

            for e in est_grado:
                s = stats_por_est.get(e.id, {"dias_presentes": 0, "ultima": None})
                dp = s["dias_presentes"] or 0
                da = max(total_dias_academia - dp, 0)
                pct = round((dp / total_dias_academia * 100), 1) if total_dias_academia > 0 else 0.0
                nombre_completo = (
                    f"{(e.apellido_paterno_est or '').strip()} "
                    f"{(e.apellido_materno_est or '').strip()}, "
                    f"{(e.nombres_est or '').strip()}"
                ).strip(" ,")
                estudiantes_academia.append(
                    {
                        "id": e.id,
                        "nombre": nombre_completo,
                        "dni": e.dni_est or "-",
                        "foto": _static_photo_url(request, e),
                        "dias_presentes": dp,
                        "dias_ausentes": da,
                        "total_dias": total_dias_academia,
                        "porcentaje": pct,
                        "ultima": s["ultima"],
                    }
                )
    else:
        query_args = {}
        if estado:
            query_args["estado"] = estado
        if fecha_inicio_obj:
            query_args["fecha_inicio"] = fecha_inicio_obj
        if fecha_fin_obj:
            query_args["fecha_fin"] = fecha_fin_obj
        if nivel and nivel_u != "ACADEMIA":
            query_args["nivel"] = nivel
            if aula_id:
                query_args["aula_id"] = int(aula_id)
        elif aula_id:
            query_args["aula_id"] = int(aula_id)

        registros = get_asistencia_service().listar_registros_aula(**query_args).all()

    # ──────────────────────────────────────────────────────────────────
    # Grid de asistencias (estilo reporte-mensual): filas=estudiantes,
    # columnas=días del rango, celda=estado (Entrada+Salida / solo ENT /
    # solo SAL / Faltó). Se activa cuando el tutor filtra por un aula
    # específica (Primaria/Secundaria) o por Academia + Programa.
    # ──────────────────────────────────────────────────────────────────
    grid_mode = False
    grid_titulo = ""
    grid_aula_obj = None
    grid_estudiantes_view: list[dict] = []
    grid_dias: list[date] = []
    grid_asistencias: dict = {}
    grid_range_inicio: date | None = None
    grid_range_fin: date | None = None
    grid_resumen = {"total": 0, "dias": 0, "presentes_avg": 0.0}

    hoy = date.today()
    _gi = fecha_inicio_obj or hoy.replace(day=1)
    _gf = fecha_fin_obj or hoy
    if _gf < _gi:
        _gi, _gf = _gf, _gi
    max_dias = 62
    if (_gf - _gi).days + 1 > max_dias:
        _gi = _gf - timedelta(days=max_dias - 1)

    estudiantes_grid_src = []
    if es_academia_grado:
        estudiantes_grid_src = est_grado if 'est_grado' in locals() else []
        grid_titulo = f"Academia / {grado}"
    elif aula_id:
        try:
            aula_id_int = int(aula_id)
        except (TypeError, ValueError):
            aula_id_int = None
        if aula_id_int:
            grid_aula_obj = db.session.get(Aula, aula_id_int)
            if grid_aula_obj:
                mats = (
                    Matricula.query.filter_by(
                        aula_id=aula_id_int, estado="activo"
                    ).all()
                )
                ids_mat = [m.estudiante_id for m in mats]
                if ids_mat:
                    estudiantes_grid_src = (
                        Estudiante.query.filter(Estudiante.id.in_(ids_mat))
                        .order_by(
                            Estudiante.apellido_paterno_est,
                            Estudiante.apellido_materno_est,
                            Estudiante.nombres_est,
                        )
                        .all()
                    )
                grid_titulo = (
                    f"{grid_aula_obj.nombre} "
                    f"({grid_aula_obj.nivel} · {grid_aula_obj.grado}"
                    f"{(' ' + grid_aula_obj.seccion) if getattr(grid_aula_obj, 'seccion', None) else ''})"
                )

    if estudiantes_grid_src:
        grid_mode = True
        grid_range_inicio = _gi
        grid_range_fin = _gf
        dia_cursor = _gi
        while dia_cursor <= _gf:
            grid_dias.append(dia_cursor)
            dia_cursor += timedelta(days=1)

        ids_est_all = [e.id for e in estudiantes_grid_src]
        asis_rows = (
            Asistencia.query.filter(
                Asistencia.estudiante_id.in_(ids_est_all),
                func.date(Asistencia.fecha_hora) >= _gi,
                func.date(Asistencia.fecha_hora) <= _gf,
            )
            .all()
        )
        for a in asis_rows:
            fecha_iso = a.fecha_hora.date().isoformat()
            key = (a.estudiante_id, fecha_iso)
            grid_asistencias.setdefault(key, []).append((a.tipo or "").upper())

        total_dias = len(grid_dias)
        suma_presencias = 0
        for e in estudiantes_grid_src:
            presentes = 0
            ausentes = 0
            solo_parcial = 0
            for d in grid_dias:
                tipos = grid_asistencias.get((e.id, d.isoformat()), [])
                if "ENTRADA" in tipos and "SALIDA" in tipos:
                    presentes += 1
                elif "ENTRADA" in tipos or "SALIDA" in tipos:
                    solo_parcial += 1
                else:
                    ausentes += 1
            suma_presencias += presentes + (solo_parcial * 0.5)
            pct = round(((presentes + solo_parcial * 0.5) / total_dias * 100), 1) if total_dias else 0.0
            nombre_completo = (
                f"{(e.apellido_paterno_est or '').strip()} "
                f"{(e.apellido_materno_est or '').strip()}, "
                f"{(e.nombres_est or '').strip()}"
            ).strip(" ,")
            grid_estudiantes_view.append(
                {
                    "id": e.id,
                    "nombre": nombre_completo,
                    "iniciales": (
                        f"{(e.nombres_est or '?')[:1]}"
                        f"{(e.apellido_paterno_est or '?')[:1]}"
                    ).upper(),
                    "foto": _static_photo_url(request, e),
                    "presentes": presentes,
                    "ausentes": ausentes,
                    "parciales": solo_parcial,
                    "porcentaje": pct,
                }
            )
        grid_resumen = {
            "total": len(estudiantes_grid_src),
            "dias": total_dias,
            "presentes_avg": round(
                (suma_presencias / len(estudiantes_grid_src)) if estudiantes_grid_src else 0,
                1,
            ),
        }

    return templates.TemplateResponse(
        "asistencias/por_aula.html",
        common_context(
            request,
            registros=registros,
            registros_academia=registros_academia,
            estudiantes_academia=estudiantes_academia,
            total_dias_academia=total_dias_academia,
            es_academia_grado=es_academia_grado,
            aulas_por_nivel=aulas_por_nivel,
            grados_academia=grados_academia,
            fecha_inicio=fecha_inicio,
            fecha_fin=fecha_fin,
            aula_id_filtro=aula_id,
            estado_filtro=estado,
            nivel_filtro=nivel,
            grado_filtro=grado,
            grid_mode=grid_mode,
            grid_titulo=grid_titulo,
            grid_dias=grid_dias,
            grid_estudiantes=grid_estudiantes_view,
            grid_asistencias=grid_asistencias,
            grid_range_inicio=grid_range_inicio,
            grid_range_fin=grid_range_fin,
            grid_resumen=grid_resumen,
        ),
    )


@router.get("/aula/{aula_id}/fecha/{fecha}", name="asistencias.tomar_asistencia")
def tomar_asistencia(
    request: Request,
    aula_id: int,
    fecha: str,
    _user_id: int = Depends(get_current_user_id),
):
    aula = db.session.get(Aula, aula_id)
    if not aula:
        add_flash(request, "Aula no encontrada", "error")
        return RedirectResponse(url=str(request.url_for("asistencias.por_aula")), status_code=303)

    try:
        fecha_obj = datetime.strptime(fecha, "%Y-%m-%d").date()
    except ValueError:
        add_flash(request, "Fecha inválida", "error")
        return RedirectResponse(url=str(request.url_for("asistencias.por_aula")), status_code=303)

    registro, error = get_asistencia_service().crear_registro_aula(
        aula_id=aula_id,
        fecha=fecha_obj,
        usuario_registro=_username(request),
    )

    if error:
        add_flash(request, error, "error")
        return RedirectResponse(url=str(request.url_for("asistencias.por_aula")), status_code=303)

    detalles = registro.detalles.order_by(DetalleAsistenciaAula.estudiante_nombre_completo).all()

    return templates.TemplateResponse(
        "asistencias/tomar_asistencia.html",
        common_context(
            request, registro=registro, aula=aula, detalles=detalles, fecha=fecha_obj
        ),
    )


@router.api_route("/aula/{aula_id}/seleccionar_fecha", methods=["GET", "POST"], name="asistencias.seleccionar_fecha_aula")
async def seleccionar_fecha_aula(
    request: Request,
    aula_id: int,
    _user_id: int = Depends(get_current_user_id),
):
    aula = db.session.get(Aula, aula_id)
    if not aula:
        add_flash(request, "Aula no encontrada", "error")
        return RedirectResponse(url=str(request.url_for("asistencias.por_aula")), status_code=303)

    if request.method == "POST":
        form = await request.form()
        if getattr(Config, "WTF_CSRF_ENABLED", True) and not csrf_ok(request, form.get("csrf_token")):
            add_flash(request, "Sesión de seguridad expirada. Intente de nuevo.", "error")
            return RedirectResponse(
                url=str(request.url_for("asistencias.seleccionar_fecha_aula", aula_id=aula_id)),
                status_code=303,
            )
        fecha = form.get("fecha")
        if not fecha:
            add_flash(request, "Debe seleccionar una fecha", "error")
            return templates.TemplateResponse(
                "asistencias/seleccionar_fecha.html", common_context(request, aula=aula)
            )

        return RedirectResponse(
            url=str(request.url_for("asistencias.tomar_asistencia", aula_id=aula_id, fecha=fecha)),
            status_code=303,
        )

    return templates.TemplateResponse(
        "asistencias/seleccionar_fecha.html", common_context(request, aula=aula)
    )


@router.post("/api/registrar_asistencia", name="asistencias.api_registrar_asistencia")
async def api_registrar_asistencia(request: Request, _user_id: int = Depends(get_current_user_id)):
    try:
        data = await request.json()
        registro_id = data.get("registro_id")
        asistencias = data.get("asistencias", {})

        if not registro_id:
            return JSONResponse({"ok": False, "message": "Falta registro_id"}, status_code=400)

        asistencias_dict = {}
        for estudiante_id, datos in asistencias.items():
            estado = datos.get("estado", "ausente")
            hora_str = datos.get("hora")
            observaciones = datos.get("observaciones")

            hora_obj = None
            if hora_str:
                try:
                    hora_obj = datetime.strptime(hora_str, "%H:%M").time()
                except ValueError:
                    pass

            asistencias_dict[int(estudiante_id)] = {
                "estado": estado,
                "hora": hora_obj,
                "observaciones": observaciones,
            }

        success, error = get_asistencia_service().registrar_asistencia_masiva(
            registro_id=registro_id,
            asistencias_dict=asistencias_dict,
            usuario=_username(request),
        )

        if not success:
            return JSONResponse({"ok": False, "message": error}, status_code=400)

        return JSONResponse({"ok": True, "message": "Asistencia registrada exitosamente"})

    except Exception as e:
        return JSONResponse({"ok": False, "message": f"Error: {str(e)}"}, status_code=500)


@router.get("/justificaciones", name="asistencias.justificaciones_lista")
def justificaciones_lista(request: Request, _user_id: int = Depends(get_current_user_id)):
    estado = request.query_params.get("estado", "")
    fecha_inicio = request.query_params.get("fecha_inicio", "")
    fecha_fin = request.query_params.get("fecha_fin", "")

    query_args = {}
    if estado:
        query_args["estado"] = estado
    if fecha_inicio:
        try:
            query_args["fecha_inicio"] = datetime.strptime(fecha_inicio, "%Y-%m-%d").date()
        except ValueError:
            pass
    if fecha_fin:
        try:
            query_args["fecha_fin"] = datetime.strptime(fecha_fin, "%Y-%m-%d").date()
        except ValueError:
            pass

    justificaciones = get_asistencia_service().listar_justificaciones(**query_args).all()

    return templates.TemplateResponse(
        "asistencias/justificaciones_lista.html",
        common_context(
            request,
            justificaciones=justificaciones,
            estado_filtro=estado,
            fecha_inicio=fecha_inicio,
            fecha_fin=fecha_fin,
        ),
    )


@router.api_route("/justificaciones/crear", methods=["GET", "POST"], name="asistencias.crear_justificacion")
async def crear_justificacion(request: Request, _user_id: int = Depends(get_current_user_id)):
    if request.method == "POST":
        form = await request.form()
        if getattr(Config, "WTF_CSRF_ENABLED", True) and not csrf_ok(request, form.get("csrf_token")):
            add_flash(request, "Sesión de seguridad expirada. Intente de nuevo.", "error")
            return templates.TemplateResponse(
                "asistencias/crear_justificacion.html", common_context(request)
            )
        estudiante_id = form.get("estudiante_id")
        fecha_str = form.get("fecha")
        apoderado_nombre = form.get("apoderado_nombre")
        apoderado_dni = form.get("apoderado_dni")
        apoderado_telefono = form.get("apoderado_telefono")
        apoderado_relacion = form.get("apoderado_relacion")
        motivo = form.get("motivo")
        tipo_motivo = form.get("tipo_motivo") or "otros"
        tiene_documento = form.get("tiene_documento") == "si"
        tipo_documento = form.get("tipo_documento") if tiene_documento else None

        if not all([estudiante_id, fecha_str, apoderado_nombre, apoderado_dni, motivo]):
            add_flash(request, "Complete todos los campos obligatorios", "error")
            return templates.TemplateResponse(
                "asistencias/crear_justificacion.html", common_context(request)
            )

        try:
            fecha = datetime.strptime(fecha_str, "%Y-%m-%d").date()
        except ValueError:
            add_flash(request, "Fecha inválida", "error")
            return templates.TemplateResponse(
                "asistencias/crear_justificacion.html", common_context(request)
            )

        justificacion, error = get_asistencia_service().crear_justificacion(
            estudiante_id=int(estudiante_id),
            fecha=fecha,
            apoderado_nombre=apoderado_nombre,
            apoderado_dni=apoderado_dni,
            apoderado_telefono=apoderado_telefono,
            apoderado_relacion=apoderado_relacion,
            motivo=motivo,
            tipo_motivo=tipo_motivo,
            tiene_documento=tiene_documento,
            tipo_documento=tipo_documento,
            usuario_registro=_username(request),
        )

        if error:
            add_flash(request, error, "error")
            return templates.TemplateResponse(
                "asistencias/crear_justificacion.html", common_context(request)
            )

        add_flash(
            request,
            f"Justificación registrada exitosamente para {justificacion.estudiante_nombre_completo}",
            "success",
        )
        return RedirectResponse(
            url=str(request.url_for("asistencias.justificaciones_lista")), status_code=303
        )

    return templates.TemplateResponse(
        "asistencias/crear_justificacion.html", common_context(request)
    )


@router.post("/api/buscar_estudiante_asistencia", name="asistencias.api_buscar_estudiante_asistencia")
async def api_buscar_estudiante_asistencia(
    request: Request, _user_id: int = Depends(get_current_user_id)
):
    try:
        try:
            data = await request.json()
        except Exception:
            data = {}
        if not isinstance(data, dict):
            data = {}
        termino = (data.get("termino") or "").strip()

        if not termino or len(termino) < 2:
            return JSONResponse(
                {"success": False, "message": "Ingrese al menos 2 caracteres"}, status_code=400
            )

        estudiante = Estudiante.query.filter_by(dni_est=termino).first()

        if not estudiante:
            estudiantes = (
                Estudiante.query.filter(
                    or_(
                        Estudiante.nombres_est.ilike(f"%{termino}%"),
                        Estudiante.apellido_paterno_est.ilike(f"%{termino}%"),
                        Estudiante.apellido_materno_est.ilike(f"%{termino}%"),
                    )
                )
                .limit(10)
                .all()
            )

            if not estudiantes:
                return JSONResponse(
                    {"success": False, "message": "No se encontraron estudiantes"}, status_code=404
                )

            if len(estudiantes) > 1:
                lista = []
                for e in estudiantes:
                    lista.append(
                        {
                            "id": e.id,
                            "nombre_completo": (
                                f"{e.apellido_paterno_est} {e.apellido_materno_est}, {e.nombres_est}"
                            ),
                            "dni": e.dni_est,
                            "nivel": e.nivel or "-",
                            "grado": e.grado or "-",
                        }
                    )
                return JSONResponse({"success": True, "multiple": True, "estudiantes": lista})

            estudiante = estudiantes[0]

        apoderado_nombre = ""
        if hasattr(estudiante, "nombres_apoderado") and estudiante.nombres_apoderado:
            partes = [estudiante.nombres_apoderado]
            if hasattr(estudiante, "apellido_paterno_apoderado") and estudiante.apellido_paterno_apoderado:
                partes.append(estudiante.apellido_paterno_apoderado)
            if hasattr(estudiante, "apellido_materno_apoderado") and estudiante.apellido_materno_apoderado:
                partes.append(estudiante.apellido_materno_apoderado)
            apoderado_nombre = " ".join(partes)

        return JSONResponse(
            {
                "success": True,
                "multiple": False,
                "estudiante": {
                    "id": estudiante.id,
                    "nombre_completo": (
                        f"{estudiante.apellido_paterno_est} {estudiante.apellido_materno_est}, "
                        f"{estudiante.nombres_est}"
                    ),
                    "dni": estudiante.dni_est,
                    "nivel": estudiante.nivel or "-",
                    "grado": estudiante.grado or "-",
                    "seccion": getattr(estudiante, "seccion", "") or "-",
                    "apoderado_nombre": apoderado_nombre,
                    "apoderado_dni": getattr(estudiante, "dni_apoderado", "") or "",
                    "apoderado_celular": getattr(estudiante, "celular_apoderado", "") or "",
                    "apoderado_relacion": getattr(estudiante, "relacion_apoderado", "") or "",
                },
            }
        )

    except Exception as e:
        return JSONResponse({"success": False, "message": f"Error: {str(e)}"}, status_code=500)


@router.get("/api/fechas_ausencia/{estudiante_id}", name="asistencias.api_fechas_ausencia")
def api_fechas_ausencia(
    request: Request,
    estudiante_id: int,
    _user_id: int = Depends(get_current_user_id),
):
    try:
        fechas = get_asistencia_service().obtener_fechas_ausencia(estudiante_id, dias=60)
        return JSONResponse({"success": True, "fechas": fechas})
    except Exception as e:
        return JSONResponse({"success": False, "message": f"Error: {str(e)}"}, status_code=500)


@router.post("/justificaciones/{justificacion_id}/aprobar", name="asistencias.aprobar_justificacion")
async def aprobar_justificacion(
    request: Request,
    justificacion_id: int,
    _user_id: int = Depends(get_current_user_id),
):
    form = await request.form()
    if getattr(Config, "WTF_CSRF_ENABLED", True) and not csrf_ok(request, form.get("csrf_token")):
        add_flash(request, "Sesión de seguridad expirada. Intente de nuevo.", "error")
        return RedirectResponse(
            url=str(request.url_for("asistencias.justificaciones_lista")), status_code=303
        )
    observaciones = form.get("observaciones")

    _j, error = get_asistencia_service().aprobar_justificacion(
        justificacion_id=justificacion_id,
        usuario=_username(request),
        observaciones=observaciones,
    )

    if error:
        add_flash(request, error, "error")
    else:
        add_flash(request, "Justificación aprobada exitosamente", "success")

    return RedirectResponse(
        url=str(request.url_for("asistencias.justificaciones_lista")), status_code=303
    )


@router.post("/justificaciones/{justificacion_id}/rechazar", name="asistencias.rechazar_justificacion")
async def rechazar_justificacion(
    request: Request,
    justificacion_id: int,
    _user_id: int = Depends(get_current_user_id),
):
    form = await request.form()
    if getattr(Config, "WTF_CSRF_ENABLED", True) and not csrf_ok(request, form.get("csrf_token")):
        add_flash(request, "Sesión de seguridad expirada. Intente de nuevo.", "error")
        return RedirectResponse(
            url=str(request.url_for("asistencias.justificaciones_lista")), status_code=303
        )
    observaciones = form.get("observaciones")

    if not observaciones:
        add_flash(request, "Debe proporcionar un motivo para rechazar", "error")
        return RedirectResponse(
            url=str(request.url_for("asistencias.justificaciones_lista")), status_code=303
        )

    _j, error = get_asistencia_service().rechazar_justificacion(
        justificacion_id=justificacion_id,
        usuario=_username(request),
        observaciones=observaciones,
    )

    if error:
        add_flash(request, error, "error")
    else:
        add_flash(request, "Justificación rechazada", "warning")

    return RedirectResponse(
        url=str(request.url_for("asistencias.justificaciones_lista")), status_code=303
    )


@router.post("/registro/{registro_id}/cerrar", name="asistencias.cerrar_registro")
async def cerrar_registro(
    request: Request,
    registro_id: int,
    _user_id: int = Depends(get_current_user_id),
):
    form = await request.form()
    if getattr(Config, "WTF_CSRF_ENABLED", True) and not csrf_ok(request, form.get("csrf_token")):
        add_flash(request, "Sesión de seguridad expirada. Intente de nuevo.", "error")
        return RedirectResponse(url=str(request.url_for("asistencias.por_aula")), status_code=303)

    _r, error = get_asistencia_service().cerrar_registro(
        registro_id=registro_id,
        usuario=_username(request),
    )

    if error:
        add_flash(request, error, "error")
    else:
        add_flash(request, "Registro de asistencia cerrado exitosamente", "success")

    return RedirectResponse(url=str(request.url_for("asistencias.por_aula")), status_code=303)
