# -*- coding: utf-8 -*-
"""
Carga de ponderaciones UNCP 2026 (80 preguntas) desde un .xlsx.

Expone funciones reutilizables tanto desde el script CLI
(``scripts/load_ponderaciones_uncp.py``) como desde el endpoint HTTP
``academia_weights``.
"""
from __future__ import annotations

import re
import unicodedata
from typing import Any, BinaryIO, Dict, List, Optional, Tuple, Union

from openpyxl import load_workbook
from sqlalchemy import or_

from models import db
from models.academia import AcademicArea, ExamenPreguntasConfig, QuestionWeight


SHEET_NAME = "PONDERADO 26"
EXPECTED_AREAS = 5
EXPECTED_QUESTIONS_PER_AREA = 80


SUBJECT_NORMALIZATION: Dict[str, str] = {
    "ARITMETICA": "Aritmética",
    "ALGEBRA": "Álgebra",
    "GEOMETRIA": "Geometría",
    "TRIGONOMETRIA": "Trigonometría",
    "ESTADISTICA Y PROBABILIDADES": "Estadística y Probabilidades",
    "COMUNICACION": "Comunicación",
    "BIOLOGIA": "Biología",
    "QUIMICA": "Química",
    "FISICA": "Física",
    "ECOLOGIA": "Ecología",
    "PSICOLOGIA": "Psicología",
    "HISTORIA": "Historia",
    "GEOGRAFIA": "Geografía",
    "ECONOMIA": "Economía",
    "CIVICA": "Cívica",
    "FILOSOFIA": "Filosofía",
    "APTITUD LOGICO MATEMATICO": "Aptitud Lógico Matemático",
    "APTITUD COMUNICATIVA": "Aptitud Comunicativa",
    # Typo en el Excel oficial UNCP 2026:
    "APTITUD COMUNICATICA": "Aptitud Comunicativa",
}

LEVEL_NORMALIZATION: Dict[str, str] = {
    "BASICO": "BÁSICO",
    "BÁSICO": "BÁSICO",
    "INTERMEDIO": "INTERMEDIO",
    "AVANZADO": "AVANZADO",
}


def _strip_accents(s: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c)
    )


def normalize_subject(raw: str) -> str:
    key = _strip_accents(raw or "").strip().upper()
    if key in SUBJECT_NORMALIZATION:
        return SUBJECT_NORMALIZATION[key]
    raise ValueError(f"Asignatura desconocida en Excel: {raw!r}")


def normalize_level(raw: str) -> str:
    key = _strip_accents(raw or "").strip().upper()
    if key in LEVEL_NORMALIZATION:
        return LEVEL_NORMALIZATION[key]
    raise ValueError(f"Nivel desconocido en Excel: {raw!r}")


def _parse_range(raw: str) -> Tuple[int, int]:
    m = re.search(r"(\d+)\s*al\s*(\d+)", str(raw or ""), flags=re.IGNORECASE)
    if not m:
        raise ValueError(f"Rango de preguntas inválido: {raw!r}")
    start, end = int(m.group(1)), int(m.group(2))
    if start < 1 or end < start:
        raise ValueError(f"Rango fuera de orden: {raw!r}")
    return start, end


def parse_workbook(source: Union[str, BinaryIO]) -> Dict[int, Dict[str, Any]]:
    """Lee el .xlsx (path o file-like) y devuelve {area_id: {'name', 'rows'}}."""
    wb = load_workbook(source, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(
            f"Hoja {SHEET_NAME!r} no encontrada. Hojas disponibles: {wb.sheetnames}"
        )
    ws = wb[SHEET_NAME]

    areas: Dict[int, Dict[str, Any]] = {}
    current_area_id: Optional[int] = None
    re_area_header = re.compile(r"AREA\s+0?(\d+)", flags=re.IGNORECASE)

    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        c = ws.cell(r, 3).value
        d = ws.cell(r, 4).value
        e = ws.cell(r, 5).value
        f = ws.cell(r, 6).value
        g = ws.cell(r, 7).value

        if isinstance(a, str):
            m = re_area_header.match(a.strip())
            if m:
                current_area_id = int(m.group(1))
                areas.setdefault(current_area_id, {"name": None, "rows": []})
                continue
            if "TOTAL" in a.strip().upper():
                current_area_id = None
                continue

        if current_area_id is None:
            continue

        if (
            areas[current_area_id]["name"] is None
            and isinstance(a, str)
            and a.strip()
        ):
            areas[current_area_id]["name"] = a.strip()

        if isinstance(d, str) and isinstance(e, str) and re.search(r"\d", d):
            try:
                q_start, q_end = _parse_range(d)
            except ValueError:
                continue
            try:
                weight = float(g)
            except (TypeError, ValueError) as exc:
                raise ValueError(
                    f"Ponderado inválido en fila {r} (col G): {g!r}"
                ) from exc
            n_preg = (q_end - q_start) + 1
            if isinstance(f, (int, float)) and int(f) != n_preg:
                raise ValueError(
                    f"AREA {current_area_id} fila {r}: 'N° PREGUNTAS'={f} "
                    f"no concuerda con rango {q_start}-{q_end} ({n_preg})."
                )
            areas[current_area_id]["rows"].append(
                {
                    "subject_raw": (c or "").strip() if isinstance(c, str) else "",
                    "level_raw": e.strip(),
                    "q_start": q_start,
                    "q_end": q_end,
                    "weight": weight,
                }
            )

    return areas


def validate_parsed(parsed: Dict[int, Dict[str, Any]]) -> None:
    if len(parsed) != EXPECTED_AREAS:
        raise ValueError(
            f"Se esperaban {EXPECTED_AREAS} áreas en el Excel, se encontraron "
            f"{len(parsed)}: {sorted(parsed.keys())}."
        )
    for area_id, info in parsed.items():
        rows_sorted = sorted(info["rows"], key=lambda x: (x["q_start"], x["q_end"]))
        expected = 1
        for r in rows_sorted:
            if r["q_start"] != expected:
                raise ValueError(
                    f"AREA {area_id}: gap/overlap en {r['subject_raw']} "
                    f"{r['q_start']}-{r['q_end']} (esperado inicio en {expected})."
                )
            expected = r["q_end"] + 1
        if expected - 1 != EXPECTED_QUESTIONS_PER_AREA:
            raise ValueError(
                f"AREA {area_id}: cubre 1..{expected-1}, se esperaba "
                f"1..{EXPECTED_QUESTIONS_PER_AREA}."
            )
        # Pre-validar que las normalizaciones existan para todas las filas.
        for r in info["rows"]:
            normalize_subject(r["subject_raw"])
            normalize_level(r["level_raw"])


def match_db_areas(parsed: Dict[int, Dict[str, Any]]) -> Dict[int, AcademicArea]:
    db_areas = {a.id: a for a in AcademicArea.query.all()}
    matched: Dict[int, AcademicArea] = {}
    for area_id, info in parsed.items():
        if area_id not in db_areas:
            raise ValueError(
                f"AREA {area_id} del Excel no existe en academic_areas (BD)."
            )
        db_area = db_areas[area_id]
        excel_tokens = set(
            _strip_accents(info["name"] or "").upper().replace(",", "").split()
        )
        db_tokens = set(
            _strip_accents(db_area.name or "").upper().replace(",", "").split()
        )
        if excel_tokens and db_tokens and len(excel_tokens & db_tokens) < 2:
            raise ValueError(
                f"Nombre del área no concuerda para AREA {area_id}: "
                f"Excel={info['name']!r} BD={db_area.name!r}."
            )
        matched[area_id] = db_area
    return matched


def apply_weights(
    parsed: Dict[int, Dict[str, Any]], commit: bool = True
) -> Dict[str, int]:
    """Reemplaza ``question_weights`` ACADEMIA para las áreas del Excel.

    Si ``commit=False`` el caller controla la transacción.
    """
    matched = match_db_areas(parsed)
    area_ids = list(matched.keys())

    deleted = (
        QuestionWeight.query.filter(
            QuestionWeight.academic_area_id.in_(area_ids),
            or_(QuestionWeight.nivel == "ACADEMIA", QuestionWeight.nivel.is_(None)),
            or_(QuestionWeight.grado.is_(None), QuestionWeight.grado == ""),
            QuestionWeight.cupo == EXPECTED_QUESTIONS_PER_AREA,
        )
        .delete(synchronize_session=False)
    )

    inserted = 0
    for area_id, info in sorted(parsed.items()):
        for row in sorted(info["rows"], key=lambda x: x["q_start"]):
            db.session.add(
                QuestionWeight(
                    academic_area_id=area_id,
                    question_start=row["q_start"],
                    question_end=row["q_end"],
                    subject=normalize_subject(row["subject_raw"]),
                    level=normalize_level(row["level_raw"]),
                    weight=row["weight"],
                    nivel="ACADEMIA",
                    grado=None,
                    cupo=EXPECTED_QUESTIONS_PER_AREA,
                )
            )
            inserted += 1
    if commit:
        db.session.commit()
    else:
        db.session.flush()
    return {"deleted": int(deleted or 0), "inserted": inserted, "areas": len(area_ids)}


def apply_cupo(
    parsed: Dict[int, Dict[str, Any]],
    target_max: int = EXPECTED_QUESTIONS_PER_AREA,
    commit: bool = True,
) -> Dict[str, int]:
    """Upsert de ``examen_preguntas_config`` (ACADEMIA, '*', area_id) = target_max."""
    matched = match_db_areas(parsed)
    upserted = 0
    for area_id in matched.keys():
        cfg = ExamenPreguntasConfig.query.filter_by(
            nivel="ACADEMIA", grado="*", academic_area_id=area_id
        ).first()
        if cfg is None:
            db.session.add(
                ExamenPreguntasConfig(
                    nivel="ACADEMIA",
                    grado="*",
                    academic_area_id=area_id,
                    max_questions=target_max,
                )
            )
        else:
            cfg.max_questions = target_max
        upserted += 1
    if commit:
        db.session.commit()
    else:
        db.session.flush()
    return {"areas": upserted, "max_questions": target_max}


def apply_full(
    parsed: Dict[int, Dict[str, Any]], commit: bool = True
) -> Dict[str, int]:
    """Aplica weights + cupo en una única transacción."""
    try:
        weights_stats = apply_weights(parsed, commit=False)
        cupo_stats = apply_cupo(parsed, commit=False)
        if commit:
            db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    return {
        "deleted_weights": weights_stats["deleted"],
        "inserted_weights": weights_stats["inserted"],
        "cupo_areas": cupo_stats["areas"],
        "cupo_max_questions": cupo_stats["max_questions"],
    }
