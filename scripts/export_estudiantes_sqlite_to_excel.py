"""
Exporta la tabla estudiantes desde instance/escuela.db a un .xlsx.

Uso:
  python scripts/export_estudiantes_sqlite_to_excel.py
  python scripts/export_estudiantes_sqlite_to_excel.py --db instance/mi_backup.db --out estudiantes.xlsx
"""
from __future__ import annotations

import argparse
import os
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Mismas columnas que la plantilla de carga masiva (routes/estudiantes.py)
PLANTILLA_HEADERS = [
    "Apellido Paterno",
    "Apellido Materno",
    "Nombres",
    "DNI",
    "Nivel",
    "Grado/Programa",
    "Carrera al que Postula",
    "Celular Estudiante",
    "Celular Padre",
    "Correo del Padre",
    "Área a la que Postula",
]


def _parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Exportar estudiantes desde SQLite a Excel.")
    p.add_argument(
        "--db",
        default=os.path.join(ROOT, "instance", "escuela.db"),
        help="Ruta al archivo SQLite (por defecto: instance/escuela.db).",
    )
    p.add_argument(
        "--out",
        default=os.path.join(ROOT, "estudiantes_desde_sqlite.xlsx"),
        help="Archivo .xlsx de salida.",
    )
    return p.parse_args()


def main() -> None:
    args = _parse_args()
    if not os.path.isfile(args.db):
        raise SystemExit(f"No existe la base: {args.db}")

    con = sqlite3.connect(args.db)
    con.row_factory = sqlite3.Row
    cur = con.cursor()
    cur.execute("SELECT * FROM estudiantes ORDER BY id")
    rows = cur.fetchall()
    col_names = [d[0] for d in cur.description] if rows else []
    if not col_names:
        cur.execute("PRAGMA table_info(estudiantes)")
        col_names = [r[1] for r in cur.fetchall()]
    con.close()

    wb = Workbook()
    # --- Hoja 1: todos los campos (backup / consulta)
    ws_full = wb.active
    ws_full.title = "Todos_los_campos"
    header_fill = PatternFill(start_color="5F2A5D", end_color="5F2A5D", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    for c, name in enumerate(col_names, 1):
        cell = ws_full.cell(row=1, column=c, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    def cell_val(row: sqlite3.Row, key: str):
        if not row:
            return None
        try:
            return row[key]
        except (KeyError, IndexError):
            return None

    for r_i, row in enumerate(rows, 2):
        for c_i, key in enumerate(col_names, 1):
            v = row[key]
            ws_full.cell(row=r_i, column=c_i, value=v)

    # --- Hoja 2: formato plantilla (reimportación masiva)
    ws_pl = wb.create_sheet(title="Plantilla_carga")
    for c, h in enumerate(PLANTILLA_HEADERS, 1):
        cell = ws_pl.cell(row=1, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    keys = {
        "apellido_paterno": "apellido_paterno_est",
        "apellido_materno": "apellido_materno_est",
        "nombres": "nombres_est",
        "dni": "dni_est",
        "nivel": "nivel",
        "grado": "grado",
        "carrera": "carrera_postula",
        "cel_est": "numero_celular_est",
        "cel_padre": "celular_padre",
        "correo_padre": "correo_padre",
        "area": "area_postula",
    }

    for r_i, row in enumerate(rows, 2):
        mapping = [
            cell_val(row, keys["apellido_paterno"]),
            cell_val(row, keys["apellido_materno"]),
            cell_val(row, keys["nombres"]),
            cell_val(row, keys["dni"]),
            cell_val(row, keys["nivel"]),
            cell_val(row, keys["grado"]),
            cell_val(row, keys["carrera"]),
            cell_val(row, keys["cel_est"]),
            cell_val(row, keys["cel_padre"]),
            cell_val(row, keys["correo_padre"]),
            cell_val(row, keys["area"]),
        ]
        for c_i, v in enumerate(mapping, 1):
            ws_pl.cell(row=r_i, column=c_i, value=v)

    wb.save(args.out)
    print(f"OK: {len(rows)} estudiantes -> {args.out}")


if __name__ == "__main__":
    main()
