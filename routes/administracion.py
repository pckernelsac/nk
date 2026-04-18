# routes/administracion.py
from __future__ import annotations

from datetime import datetime, timedelta
from io import BytesIO

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import RedirectResponse, Response
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func, or_

from config import Config
from dependencies import get_current_user_id, require_roles
from models import AuditLog, ConfiguracionSistema, Usuario, db
from template_helpers import add_flash, common_context, csrf_ok, templates
from utils.audit import registrar_auditoria

router = APIRouter()


@router.get("/dashboard", name="administracion.dashboard")
def dashboard(
    request: Request,
    _uid: int = Depends(get_current_user_id),
    _role: int = Depends(require_roles("administrador")),
):
    try:
        total_usuarios = Usuario.query.count()
        usuarios_activos = Usuario.query.filter_by(activo=True).count()
        usuarios_inactivos = total_usuarios - usuarios_activos

        usuarios_por_rol = (
            db.session.query(Usuario.rol, func.count(Usuario.id))
            .group_by(Usuario.rol)
            .all()
        )

        logs_recientes = (
            AuditLog.query.order_by(AuditLog.fecha_hora.desc()).limit(10).all()
        )
        usuarios_recientes = (
            Usuario.query.order_by(Usuario.fecha_registro.desc()).limit(5).all()
        )
        config = ConfiguracionSistema.query.filter_by(activo=True).first()

        stats = {
            "total_usuarios": total_usuarios,
            "usuarios_activos": usuarios_activos,
            "usuarios_inactivos": usuarios_inactivos,
            "usuarios_por_rol": dict(usuarios_por_rol),
            "total_logs": AuditLog.query.count(),
        }

        return templates.TemplateResponse(
            "administracion/dashboard.html",
            common_context(
                request,
                stats=stats,
                logs_recientes=logs_recientes,
                usuarios_recientes=usuarios_recientes,
                config=config,
            ),
        )
    except Exception as e:
        print(f"Error en dashboard: {e}")
        add_flash(request, "Error al cargar el dashboard", "error")
        return templates.TemplateResponse(
            "administracion/dashboard.html",
            common_context(
                request,
                stats={},
                logs_recientes=[],
                usuarios_recientes=[],
                config=None,
            ),
        )


@router.get("/usuarios", name="administracion.usuarios_lista")
def usuarios_lista(
    request: Request,
    _uid: int = Depends(get_current_user_id),
    _role: int = Depends(require_roles("administrador")),
):
    try:
        q = (request.query_params.get("q") or "").strip()
        rol = (request.query_params.get("rol") or "").strip()
        estado = (request.query_params.get("estado") or "").strip()

        query = Usuario.query

        if q:
            query = query.filter(
                or_(
                    Usuario.username.ilike(f"%{q}%"),
                    Usuario.nombre.ilike(f"%{q}%"),
                    Usuario.email.ilike(f"%{q}%"),
                )
            )

        if rol:
            query = query.filter(Usuario.rol == rol)

        if estado == "activo":
            query = query.filter(Usuario.activo == True)  # noqa: E712
        elif estado == "inactivo":
            query = query.filter(Usuario.activo == False)  # noqa: E712

        usuarios = query.order_by(Usuario.fecha_registro.desc()).all()
        roles_disponibles = ["administrador", "docente", "auxiliar", "contador"]

        return templates.TemplateResponse(
            "administracion/usuarios_lista.html",
            common_context(
                request,
                usuarios=usuarios,
                roles_disponibles=roles_disponibles,
                filtros={"q": q, "rol": rol, "estado": estado},
            ),
        )
    except Exception as e:
        print(f"Error al cargar usuarios: {e}")
        add_flash(request, "Error al cargar la lista de usuarios", "error")
        return templates.TemplateResponse(
            "administracion/usuarios_lista.html",
            common_context(request, usuarios=[], roles_disponibles=[]),
        )


async def _usuarios_form_handler(
    request: Request, usuario_id: int | None, _uid: int, _role: int
):
    usuario = None
    es_edicion = usuario_id is not None

    if es_edicion:
        usuario = db.session.get(Usuario, usuario_id)
        if usuario is None:
            raise HTTPException(status_code=404)
        if usuario.id == request.session.get("user_id"):
            add_flash(
                request,
                "No puedes modificar tu propio usuario desde aquí",
                "warning",
            )
            return RedirectResponse(
                url=str(request.url_for("administracion.usuarios_lista")), status_code=303
            )

    if request.method == "POST":
        form = await request.form()
        if getattr(Config, "WTF_CSRF_ENABLED", True) and not csrf_ok(request, form.get("csrf_token")):
            add_flash(request, "Sesión de seguridad expirada. Intente de nuevo.", "error")
            return RedirectResponse(str(request.url), status_code=303)
        try:
            username = (form.get("username") or "").strip()
            email = (form.get("email") or "").strip()
            nombre = (form.get("nombre") or "").strip()
            rol = (form.get("rol") or "usuario").strip()
            activo = form.get("activo") == "on"
            password = (form.get("password") or "").strip()

            if not username or not email or not nombre:
                add_flash(request, "Todos los campos son obligatorios", "error")
                return RedirectResponse(str(request.url), status_code=303)

            if not es_edicion or usuario.username != username:
                if Usuario.query.filter_by(username=username).first():
                    add_flash(request, "El nombre de usuario ya existe", "error")
                    return RedirectResponse(str(request.url), status_code=303)

            if not es_edicion or usuario.email != email:
                if Usuario.query.filter_by(email=email).first():
                    add_flash(request, "El correo electrónico ya existe", "error")
                    return RedirectResponse(str(request.url), status_code=303)

            if es_edicion:
                datos_anteriores = {
                    "username": usuario.username,
                    "email": usuario.email,
                    "nombre": usuario.nombre,
                    "rol": usuario.rol,
                    "activo": usuario.activo,
                }

                if usuario.rol == "administrador" and usuario.activo and not activo:
                    admins_activos = Usuario.query.filter_by(
                        rol="administrador", activo=True
                    ).count()
                    if admins_activos <= 1:
                        add_flash(
                            request,
                            "No puedes desactivar al único administrador activo del sistema",
                            "error",
                        )
                        return RedirectResponse(str(request.url), status_code=303)

                usuario.username = username
                usuario.email = email
                usuario.nombre = nombre
                usuario.rol = rol
                usuario.activo = activo

                if password:
                    usuario.set_password(password)

                datos_nuevos = {
                    "username": username,
                    "email": email,
                    "nombre": nombre,
                    "rol": rol,
                    "activo": activo,
                }

                db.session.commit()

                registrar_auditoria(
                    accion="UPDATE",
                    modulo="usuarios",
                    descripcion=f"Usuario actualizado: {username}",
                    entidad_tipo="Usuario",
                    entidad_id=usuario.id,
                    datos_anteriores=datos_anteriores,
                    datos_nuevos=datos_nuevos,
                )

                add_flash(request, f"Usuario {username} actualizado exitosamente", "success")

            else:
                if not password:
                    add_flash(
                        request,
                        "La contraseña es obligatoria para nuevos usuarios",
                        "error",
                    )
                    return RedirectResponse(str(request.url), status_code=303)

                nuevo_usuario = Usuario(
                    username=username,
                    email=email,
                    nombre=nombre,
                    rol=rol,
                    activo=activo,
                )
                nuevo_usuario.set_password(password)

                db.session.add(nuevo_usuario)
                db.session.commit()

                registrar_auditoria(
                    accion="CREATE",
                    modulo="usuarios",
                    descripcion=f"Usuario creado: {username}",
                    entidad_tipo="Usuario",
                    entidad_id=nuevo_usuario.id,
                    datos_nuevos={
                        "username": username,
                        "email": email,
                        "nombre": nombre,
                        "rol": rol,
                        "activo": activo,
                    },
                )

                add_flash(request, f"Usuario {username} creado exitosamente", "success")

            return RedirectResponse(
                url=str(request.url_for("administracion.usuarios_lista")), status_code=303
            )

        except Exception as e:
            db.session.rollback()
            print(f"Error al guardar usuario: {e}")
            add_flash(request, "Error al guardar el usuario. Intente nuevamente.", "error")

    roles_disponibles = ["administrador", "docente", "auxiliar", "contador"]
    return templates.TemplateResponse(
        "administracion/usuarios_form.html",
        common_context(
            request,
            usuario=usuario,
            es_edicion=es_edicion,
            roles_disponibles=roles_disponibles,
        ),
    )


@router.api_route(
    "/usuarios/crear",
    methods=["GET", "POST"],
    name="administracion.usuarios_form_crear",
)
async def usuarios_form_crear(
    request: Request,
    _uid: int = Depends(get_current_user_id),
    _role: int = Depends(require_roles("administrador")),
):
    return await _usuarios_form_handler(request, None, _uid, _role)


@router.api_route(
    "/usuarios/{usuario_id}/editar",
    methods=["GET", "POST"],
    name="administracion.usuarios_form_editar",
)
async def usuarios_form_editar(
    request: Request,
    usuario_id: int,
    _uid: int = Depends(get_current_user_id),
    _role: int = Depends(require_roles("administrador")),
):
    return await _usuarios_form_handler(request, usuario_id, _uid, _role)


@router.post("/usuarios/{usuario_id}/toggle", name="administracion.usuarios_toggle")
async def usuarios_toggle(
    request: Request,
    usuario_id: int,
    _uid: int = Depends(get_current_user_id),
    _role: int = Depends(require_roles("administrador")),
):
    form = await request.form()
    if getattr(Config, "WTF_CSRF_ENABLED", True) and not csrf_ok(request, form.get("csrf_token")):
        add_flash(request, "Sesión de seguridad expirada. Intente de nuevo.", "error")
        return RedirectResponse(
            url=str(request.url_for("administracion.usuarios_lista")), status_code=303
        )
    try:
        usuario = db.session.get(Usuario, usuario_id)
        if usuario is None:
            raise HTTPException(status_code=404)

        if usuario.id == request.session.get("user_id"):
            add_flash(request, "No puedes desactivar tu propio usuario", "error")
            return RedirectResponse(
                url=str(request.url_for("administracion.usuarios_lista")), status_code=303
            )

        if usuario.rol == "administrador" and usuario.activo:
            admins_activos = Usuario.query.filter_by(
                rol="administrador", activo=True
            ).count()
            if admins_activos <= 1:
                add_flash(
                    request,
                    "No puedes desactivar al único administrador activo del sistema",
                    "error",
                )
                return RedirectResponse(
                    url=str(request.url_for("administracion.usuarios_lista")), status_code=303
                )

        estado_anterior = usuario.activo
        usuario.activo = not usuario.activo

        db.session.commit()

        accion_texto = "activado" if usuario.activo else "desactivado"
        registrar_auditoria(
            accion="UPDATE",
            modulo="usuarios",
            descripcion=f"Usuario {accion_texto}: {usuario.username}",
            entidad_tipo="Usuario",
            entidad_id=usuario.id,
            datos_anteriores={"activo": estado_anterior},
            datos_nuevos={"activo": usuario.activo},
        )

        add_flash(
            request,
            f"Usuario {usuario.username} {accion_texto} exitosamente",
            "success",
        )

    except Exception as e:
        db.session.rollback()
        print(f"Error al cambiar estado: {e}")
        add_flash(request, "Error al cambiar el estado del usuario", "error")

    return RedirectResponse(
        url=str(request.url_for("administracion.usuarios_lista")), status_code=303
    )


@router.api_route(
    "/usuarios/{usuario_id}/cambiar-password",
    methods=["GET", "POST"],
    name="administracion.usuarios_cambiar_password",
)
async def usuarios_cambiar_password(
    request: Request,
    usuario_id: int,
    _uid: int = Depends(get_current_user_id),
    _role: int = Depends(require_roles("administrador")),
):
    usuario = db.session.get(Usuario, usuario_id)
    if usuario is None:
        raise HTTPException(status_code=404)

    if request.method == "POST":
        form = await request.form()
        if getattr(Config, "WTF_CSRF_ENABLED", True) and not csrf_ok(request, form.get("csrf_token")):
            add_flash(request, "Sesión de seguridad expirada. Intente de nuevo.", "error")
            return RedirectResponse(str(request.url), status_code=303)
        try:
            nueva_password = (form.get("nueva_password") or "").strip()
            confirmar_password = (form.get("confirmar_password") or "").strip()

            if not nueva_password or not confirmar_password:
                add_flash(request, "Todos los campos son obligatorios", "error")
                return RedirectResponse(str(request.url), status_code=303)

            if len(nueva_password) < 6:
                add_flash(
                    request,
                    "La contraseña debe tener al menos 6 caracteres",
                    "error",
                )
                return RedirectResponse(str(request.url), status_code=303)

            if nueva_password != confirmar_password:
                add_flash(request, "Las contraseñas no coinciden", "error")
                return RedirectResponse(str(request.url), status_code=303)

            usuario.set_password(nueva_password)
            db.session.commit()

            registrar_auditoria(
                accion="UPDATE",
                modulo="usuarios",
                descripcion=f"Contraseña actualizada para usuario: {usuario.username}",
                entidad_tipo="Usuario",
                entidad_id=usuario.id,
            )

            add_flash(
                request,
                f"Contraseña actualizada exitosamente para {usuario.username}",
                "success",
            )
            return RedirectResponse(
                url=str(request.url_for("administracion.usuarios_lista")), status_code=303
            )

        except Exception as e:
            db.session.rollback()
            print(f"Error al cambiar contraseña: {e}")
            add_flash(
                request,
                "Error al cambiar la contraseña. Intente nuevamente.",
                "error",
            )

    return templates.TemplateResponse(
        "administracion/usuarios_cambiar_password.html",
        common_context(request, usuario=usuario),
    )


@router.get("/roles", name="administracion.roles_info")
def roles_info(
    request: Request,
    _uid: int = Depends(get_current_user_id),
    _role: int = Depends(require_roles("administrador")),
):
    permisos = {
        "administrador": {
            "nombre": "Administrador",
            "descripcion": "Acceso total al sistema",
            "permisos": [
                "Gestión completa de usuarios",
                "Configuración del sistema",
                "Acceso a todos los módulos",
                "Gestión de estudiantes y docentes",
                "Control de pensiones",
                "Control de asistencias",
                "Auditoría del sistema",
            ],
        },
        "docente": {
            "nombre": "Docente",
            "descripcion": "Profesor del colegio",
            "permisos": [
                "Ver lista de estudiantes",
                "Ver lista de docentes",
                "Consultar asistencias",
                "Acceso de solo lectura",
            ],
        },
        "auxiliar": {
            "nombre": "Auxiliar",
            "descripcion": "Personal de apoyo",
            "permisos": [
                "Registro de asistencias (escaneo)",
                "Dashboard de asistencias",
                "Reportes de asistencias",
                "Ver lista de estudiantes",
            ],
        },
        "contador": {
            "nombre": "Contador/Tesorero",
            "descripcion": "Gestión financiera",
            "permisos": [
                "Gestión completa de pensiones",
                "Registro de pagos",
                "Generación de recibos",
                "Reportes financieros",
                "Ver lista de estudiantes",
            ],
        },
    }

    return templates.TemplateResponse(
        "administracion/roles_info.html", common_context(request, permisos=permisos)
    )


@router.api_route(
    "/configuracion",
    methods=["GET", "POST"],
    name="administracion.configuracion",
)
async def configuracion(
    request: Request,
    _uid: int = Depends(get_current_user_id),
    _role: int = Depends(require_roles("administrador")),
):
    if request.method == "POST":
        form = await request.form()
        if getattr(Config, "WTF_CSRF_ENABLED", True) and not csrf_ok(request, form.get("csrf_token")):
            add_flash(request, "Sesión de seguridad expirada. Intente de nuevo.", "error")
            return RedirectResponse(
                url=str(request.url_for("administracion.configuracion")), status_code=303
            )
        try:
            config = ConfiguracionSistema.query.filter_by(activo=True).first()

            datos_form = {
                "nombre_institucion": (form.get("nombre_institucion") or "").strip(),
                "ruc": (form.get("ruc") or "").strip(),
                "direccion": (form.get("direccion") or "").strip(),
                "telefono": (form.get("telefono") or "").strip(),
                "email": (form.get("email") or "").strip(),
                "sitio_web": (form.get("sitio_web") or "").strip(),
                "anio_academico_actual": (form.get("anio_academico_actual") or "").strip(),
                "zona_horaria": (form.get("zona_horaria") or "America/Lima").strip(),
                "idioma": (form.get("idioma") or "es").strip(),
                "moneda": (form.get("moneda") or "PEN").strip(),
            }

            if config:
                datos_anteriores = {
                    "nombre_institucion": config.nombre_institucion,
                    "anio_academico_actual": config.anio_academico_actual,
                }

                for key, value in datos_form.items():
                    setattr(config, key, value)

                config.fecha_modificacion = datetime.utcnow()
                config.usuario_modificacion = request.session.get("username")

                mensaje = "Configuración actualizada exitosamente"

            else:
                config = ConfiguracionSistema(**datos_form)
                config.activo = True
                config.usuario_modificacion = request.session.get("username")
                db.session.add(config)

                datos_anteriores = None
                mensaje = "Configuración creada exitosamente"

            db.session.commit()

            registrar_auditoria(
                accion="UPDATE" if datos_anteriores else "CREATE",
                modulo="configuracion",
                descripcion=mensaje,
                entidad_tipo="ConfiguracionSistema",
                entidad_id=config.id,
                datos_anteriores=datos_anteriores,
                datos_nuevos=datos_form,
            )

            add_flash(request, mensaje, "success")
            return RedirectResponse(
                url=str(request.url_for("administracion.configuracion")), status_code=303
            )

        except Exception as e:
            db.session.rollback()
            print(f"Error al guardar configuración: {e}")
            add_flash(request, "Error al guardar la configuración", "error")

    config = ConfiguracionSistema.query.filter_by(activo=True).first()
    return templates.TemplateResponse(
        "administracion/configuracion.html", common_context(request, config=config)
    )


@router.get("/auditoria", name="administracion.auditoria_lista")
def auditoria_lista(
    request: Request,
    _uid: int = Depends(get_current_user_id),
    _role: int = Depends(require_roles("administrador")),
):
    try:
        usuario = (request.query_params.get("usuario") or "").strip()
        accion = (request.query_params.get("accion") or "").strip()
        modulo = (request.query_params.get("modulo") or "").strip()
        fecha_desde = (request.query_params.get("fecha_desde") or "").strip()
        fecha_hasta = (request.query_params.get("fecha_hasta") or "").strip()

        query = AuditLog.query

        if usuario:
            query = query.filter(AuditLog.usuario_username.ilike(f"%{usuario}%"))

        if accion:
            query = query.filter(AuditLog.accion == accion)

        if modulo:
            query = query.filter(AuditLog.modulo == modulo)

        if fecha_desde:
            query = query.filter(AuditLog.fecha_hora >= fecha_desde)

        if fecha_hasta:
            fecha_hasta_dt = datetime.strptime(fecha_hasta, "%Y-%m-%d") + timedelta(days=1)
            query = query.filter(AuditLog.fecha_hora < fecha_hasta_dt)

        logs = query.order_by(AuditLog.fecha_hora.desc()).limit(500).all()

        acciones_disponibles = [
            "CREATE",
            "UPDATE",
            "DELETE",
            "LOGIN",
            "LOGOUT",
            "VIEW",
        ]
        modulos_disponibles = [
            "usuarios",
            "estudiantes",
            "docentes",
            "pensiones",
            "asistencias",
            "configuracion",
            "auth",
        ]

        return templates.TemplateResponse(
            "administracion/auditoria_lista.html",
            common_context(
                request,
                logs=logs,
                acciones_disponibles=acciones_disponibles,
                modulos_disponibles=modulos_disponibles,
                filtros={
                    "usuario": usuario,
                    "accion": accion,
                    "modulo": modulo,
                    "fecha_desde": fecha_desde,
                    "fecha_hasta": fecha_hasta,
                },
            ),
        )
    except Exception as e:
        print(f"Error al cargar logs: {e}")
        add_flash(request, "Error al cargar los registros de auditoría", "error")
        return templates.TemplateResponse(
            "administracion/auditoria_lista.html",
            common_context(
                request,
                logs=[],
                acciones_disponibles=[],
                modulos_disponibles=[],
            ),
        )


@router.get("/auditoria/export", name="administracion.auditoria_export")
def auditoria_export(
    request: Request,
    _uid: int = Depends(get_current_user_id),
    _role: int = Depends(require_roles("administrador")),
):
    try:
        usuario = (request.query_params.get("usuario") or "").strip()
        accion = (request.query_params.get("accion") or "").strip()
        modulo = (request.query_params.get("modulo") or "").strip()
        fecha_desde = (request.query_params.get("fecha_desde") or "").strip()
        fecha_hasta = (request.query_params.get("fecha_hasta") or "").strip()

        query = AuditLog.query

        if usuario:
            query = query.filter(AuditLog.usuario_username.ilike(f"%{usuario}%"))
        if accion:
            query = query.filter(AuditLog.accion == accion)
        if modulo:
            query = query.filter(AuditLog.modulo == modulo)
        if fecha_desde:
            query = query.filter(AuditLog.fecha_hora >= fecha_desde)
        if fecha_hasta:
            fecha_hasta_dt = datetime.strptime(fecha_hasta, "%Y-%m-%d") + timedelta(days=1)
            query = query.filter(AuditLog.fecha_hora < fecha_hasta_dt)

        logs = query.order_by(AuditLog.fecha_hora.desc()).all()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Auditoría"

        headers = [
            "ID",
            "Usuario",
            "Rol",
            "Acción",
            "Módulo",
            "Descripción",
            "Entidad",
            "IP",
            "Fecha y Hora",
        ]
        ws.append(headers)

        header_fill = PatternFill(
            start_color="4472C4", end_color="4472C4", fill_type="solid"
        )
        header_font = Font(bold=True, color="FFFFFF")

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for log in logs:
            entidad_info = (
                f"{log.entidad_tipo} #{log.entidad_id}"
                if log.entidad_tipo
                else "-"
            )
            ws.append(
                [
                    log.id,
                    log.usuario_username,
                    log.usuario_rol or "-",
                    log.accion,
                    log.modulo,
                    log.descripcion or "-",
                    entidad_info,
                    log.ip_address or "-",
                    log.fecha_hora.strftime("%Y-%m-%d %H:%M:%S"),
                ]
            )

        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 15
        ws.column_dimensions["C"].width = 15
        ws.column_dimensions["D"].width = 12
        ws.column_dimensions["E"].width = 15
        ws.column_dimensions["F"].width = 40
        ws.column_dimensions["G"].width = 20
        ws.column_dimensions["H"].width = 15
        ws.column_dimensions["I"].width = 20

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"auditoria_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename={filename}",
                "Cache-Control": "no-cache, no-store, must-revalidate",
                "Pragma": "no-cache",
                "Expires": "0",
            },
        )

    except Exception as e:
        print(f"Error al exportar logs: {e}")
        add_flash(request, "Error al exportar los registros de auditoría", "error")
        return RedirectResponse(
            url=str(request.url_for("administracion.auditoria_lista")), status_code=303
        )
