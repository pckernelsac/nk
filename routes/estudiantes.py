# routes/estudiantes.py
from __future__ import annotations

import os
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from urllib.parse import urlencode

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile
from fastapi.responses import JSONResponse, RedirectResponse, Response
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import extract, func, or_
from config import Config
from dependencies import get_current_user_id
from models import Asistencia, Estudiante, db
from models.aula import Aula
from models.matricula import Matricula
from services.acceso_portal import anio_academico_actual, ids_sin_matricula_vigente
from services.aula_service import AulaService
from template_helpers import add_flash, common_context, csrf_ok, templates
from utils.carnet_generator import generar_carnet_imagen, generar_carnets_a4
from utils.pagination import paginate_query

router = APIRouter()
ROOT_DIR = Path(__file__).resolve().parents[1]


def _qp_int(qp, key: str, default: int | None = None) -> int | None:
    v = qp.get(key)
    if v is None or v == "":
        return default
    try:
        return int(v)
    except ValueError:
        return default


def _excel_upload_from_form(form) -> UploadFile | None:
    """Obtiene el archivo Excel del multipart (nombre de campo puede variar)."""
    for key in ("excel_file", "file", "excel"):
        f = form.get(key)
        if f is None:
            continue
        if isinstance(f, UploadFile):
            return f
        # Fallback si el tipo no coincide (versiones distintas de Starlette/multipart)
        if getattr(f, "filename", None) is not None and callable(getattr(f, "read", None)):
            return f  # type: ignore[return-value]
    return None


# Tokens que tratamos como "vacío" al leer celdas de Excel. Cubren el caso de
# planillas exportadas con pandas/openpyxl donde NaN/None se serializan como
# texto y luego se reimportan, contaminando la BD con strings literales.
_EMPTY_CELL_TOKENS = {"", "none", "null", "nan", "n/a", "na", "#n/a"}


def _clean_cell(value) -> str | None:
    """Normaliza una celda: devuelve None si está vacía o contiene un token
    placeholder ('None', 'NULL', 'nan', 'N/A', etc., case-insensitive)."""
    if value is None:
        return None
    s = str(value).strip()
    if s.lower() in _EMPTY_CELL_TOKENS:
        return None
    return s


def _grados_activos_por_nivel() -> dict[str, list[str]]:
    """{NIVEL_UPPER: [grado1, grado2, ...]} desde aulas activas.

    Fuente única de verdad para los dropdowns de Grado/Programa en los
    formularios de estudiantes — solo aparecen valores que existen en /aulas.
    """
    out: dict[str, list[str]] = {}
    try:
        filas = (
            db.session.query(Aula.nivel, Aula.grado)
            .filter_by(activo=True)
            .order_by(Aula.nivel, Aula.grado)
            .distinct()
            .all()
        )
        for nivel, grado in filas:
            if not nivel or not grado:
                continue
            key = nivel.upper()
            out.setdefault(key, [])
            if grado not in out[key]:
                out[key].append(grado)
    except Exception:
        pass
    return out


def allowed_file(filename: str | None) -> bool:
    if not filename or "." not in filename:
        return False
    ext = filename.rsplit(".", 1)[1].lower()
    return ext in Config.ALLOWED_EXTENSIONS


async def save_student_photo_upload(upload: UploadFile | None, student_dni: str | None) -> str | None:
    if not upload or not upload.filename or not allowed_file(upload.filename):
        return None
    ext = upload.filename.rsplit(".", 1)[1].lower()
    filename = f"estudiante_{student_dni}_{datetime.now().strftime('%Y%m%d%H%M%S')}.{ext}"
    upload_folder = Config.STUDENT_PHOTOS_FOLDER
    os.makedirs(upload_folder, exist_ok=True)
    filepath = os.path.join(upload_folder, filename)
    data = await upload.read()
    with open(filepath, "wb") as f:
        f.write(data)
    return f"uploads/estudiantes/{filename}"


def _get_estudiante_or_404(estudiante_id: int) -> Estudiante:
    e = db.session.get(Estudiante, estudiante_id)
    if e is None:
        raise HTTPException(status_code=404)
    return e


def _query_sin_matricula_vigente(query, anio: str):
    """Restringe la consulta a quienes tienen matrícula pero ninguna del ciclo."""
    con_alguna = db.session.query(Matricula.estudiante_id).distinct()
    vigentes = (
        db.session.query(Matricula.estudiante_id)
        .filter(Matricula.estado == "activo", Matricula.anio_escolar == anio)
        .distinct()
    )
    return query.filter(
        Estudiante.id.in_(con_alguna.subquery().select()),
        ~Estudiante.id.in_(vigentes.subquery().select()),
    )


@router.get("/estudiantes", name="estudiantes.list")
def list_estudiantes(request: Request, _user_id: int = Depends(get_current_user_id)):
    """Lista de estudiantes con paginación y búsqueda"""
    page = int(request.query_params.get("page") or 1)
    per_page = 20
    search = (request.query_params.get("q") or "").strip()
    filtro = (request.query_params.get("filtro") or "").strip()
    anio_ciclo = anio_academico_actual()

    query = Estudiante.query

    if filtro == "sin_matricula":
        query = _query_sin_matricula_vigente(query, anio_ciclo)

    if search:
        search_filter = f"%{search}%"
        query = query.filter(
            or_(
                Estudiante.nombres_est.ilike(search_filter),
                Estudiante.apellido_paterno_est.ilike(search_filter),
                Estudiante.apellido_materno_est.ilike(search_filter),
                Estudiante.dni_est.ilike(search_filter),
                Estudiante.grado.ilike(search_filter),
                Estudiante.nivel.ilike(search_filter),
            )
        )

    query = query.order_by(Estudiante.apellido_paterno_est, Estudiante.nombres_est)
    estudiantes = paginate_query(query, page=page, per_page=per_page, error_out=False)

    # Marca en la página actual a quién le tocaría quedarse fuera del portal, y
    # cuántos son en toda la base (para el botón de suspensión en lote).
    try:
        sin_matricula_ids = ids_sin_matricula_vigente(
            [e.id for e in estudiantes.items], anio_ciclo
        )
        total_sin_matricula = _query_sin_matricula_vigente(
            db.session.query(Estudiante.id), anio_ciclo
        ).count()
    except Exception as e:  # noqa: BLE001
        print(f"Error al calcular matrículas vigentes: {e}")
        sin_matricula_ids = set()
        total_sin_matricula = 0

    return templates.TemplateResponse(
        "estudiantes.html",
        common_context(
            request,
            estudiantes=estudiantes,
            search=search,
            filtro=filtro,
            anio_ciclo=anio_ciclo,
            sin_matricula_ids=sin_matricula_ids,
            total_sin_matricula=total_sin_matricula,
        ),
    )


@router.get("/estudiante/{id}", name="estudiantes.ver")
def ver(request: Request, id: int, _user_id: int = Depends(get_current_user_id)):
    """Ver detalle de un estudiante"""
    estudiante = _get_estudiante_or_404(id)

    from models.academia import AcademiaStudent, AcademicArea

    etas_por_programa = {}
    if estudiante.dni_est:
        etas_raw = (
            db.session.query(AcademiaStudent, AcademicArea.name.label("area_nombre"))
            .outerjoin(AcademicArea, AcademiaStudent.academic_area_id == AcademicArea.id)
            .filter(
                or_(
                    AcademiaStudent.estudiante_id == estudiante.id,
                    AcademiaStudent.student_id == estudiante.dni_est,
                )
            )
            .order_by(AcademiaStudent.programa, AcademiaStudent.quiz_created.asc())
            .all()
        )

        for eta, area_nombre in etas_raw:
            programa = eta.programa or "Sin programa"
            if programa not in etas_por_programa:
                etas_por_programa[programa] = {"area": area_nombre or "Sin área", "etas": []}
            etas_por_programa[programa]["etas"].append(eta)

    from models.academico import FastTest, NotaFastTest

    fast_tests_por_programa = {}
    notas_ft = (
        NotaFastTest.query.join(FastTest)
        .filter(NotaFastTest.estudiante_id == id)
        .order_by(FastTest.aula_nombre, FastTest.fecha_evaluacion.asc())
        .all()
    )

    for nota in notas_ft:
        ft = nota.fast_test
        programa = ft.aula_nombre or "Sin aula"
        if programa not in fast_tests_por_programa:
            fast_tests_por_programa[programa] = []
        fast_tests_por_programa[programa].append(
            {
                "titulo": ft.titulo,
                "curso": ft.curso_nombre,
                "fecha": ft.fecha_evaluacion,
                "nota": nota.nota,
                "nota_maxima": ft.nota_maxima,
                "estado": ft.estado,
            }
        )

    return templates.TemplateResponse(
        "ver_estudiante.html",
        common_context(
            request,
            estudiante=estudiante,
            etas_por_programa=etas_por_programa,
            fast_tests_por_programa=fast_tests_por_programa,
        ),
    )


@router.api_route("/estudiante/{id}/editar", methods=["GET", "POST"], name="estudiantes.editar")
async def editar(request: Request, id: int, _user_id: int = Depends(get_current_user_id)):
    """Editar estudiante"""
    estudiante = _get_estudiante_or_404(id)

    if request.method == "POST":
        form = await request.form()
        if getattr(Config, "WTF_CSRF_ENABLED", True) and not csrf_ok(request, form.get("csrf_token")):
            add_flash(request, "Sesión de seguridad expirada. Intente de nuevo.", "error")
            return RedirectResponse(url=str(request.url_for("estudiantes.editar", id=id)), status_code=303)
        try:
            estudiante.nivel = form.get("nivel")
            estudiante.grado = form.get("grado")
            estudiante.seccion = form.get("seccion")
            estudiante.turno = form.get("turno")
            estudiante.area_postula = form.get("area_postula")
            estudiante.apellido_paterno_est = form.get("apellido_paterno_est")
            estudiante.apellido_materno_est = form.get("apellido_materno_est")
            estudiante.nombres_est = form.get("nombres_est")
            estudiante.dni_est = form.get("dni_est")
            estudiante.correo_est = form.get("correo_est")
            estudiante.fecha_nacimiento_est = form.get("fecha_nacimiento_est")
            estudiante.direccion_est = form.get("direccion_est")
            estudiante.distrito_est = form.get("distrito_est")
            estudiante.provincia_est = form.get("provincia_est")
            estudiante.referencia_est = form.get("referencia_est")
            estudiante.numero_celular_est = form.get("numero_celular_est")
            estudiante.telefono_fijo_est = form.get("telefono_fijo_est")
            estudiante.religion_est = form.get("religion_est")
            estudiante.tiene_hermanos = form.get("tiene_hermanos")
            estudiante.numero_hermanos = form.get("numero_hermanos")
            estudiante.tiene_computadora = form.get("tiene_computadora")
            estudiante.grupo_sanguineo = form.get("grupo_sanguineo")
            estudiante.es_alergica = form.get("es_alergica")
            estudiante.padece_enfermedad = form.get("padece_enfermedad")
            estudiante.tiene_discapacidad = form.get("tiene_discapacidad")

            estudiante.apellido_paterno_padre = form.get("apellido_paterno_padre")
            estudiante.apellido_materno_padre = form.get("apellido_materno_padre")
            estudiante.nombres_padre = form.get("nombres_padre")
            estudiante.dni_padre = form.get("dni_padre")
            estudiante.grado_instruccion_padre = form.get("grado_instruccion_padre")
            estudiante.fecha_nacimiento_padre = form.get("fecha_nacimiento_padre")
            estudiante.direccion_padre = form.get("direccion_padre")
            estudiante.distrito_padre = form.get("distrito_padre")
            estudiante.provincia_padre = form.get("provincia_padre")
            estudiante.celular_padre = form.get("celular_padre")
            estudiante.telefono_fijo_padre = form.get("telefono_fijo_padre")
            estudiante.ocupacion_padre = form.get("ocupacion_padre")
            estudiante.centro_trabajo_padre = form.get("centro_trabajo_padre")
            estudiante.telefono_trabajo_padre = form.get("telefono_trabajo_padre")
            estudiante.estado_civil_padre = form.get("estado_civil_padre")
            estudiante.religion_padre = form.get("religion_padre")
            estudiante.vive_con_hijo_padre = form.get("vive_con_hijo_padre")
            estudiante.correo_padre = form.get("correo_padre")

            estudiante.apellido_paterno_madre = form.get("apellido_paterno_madre")
            estudiante.apellido_materno_madre = form.get("apellido_materno_madre")
            estudiante.nombres_madre = form.get("nombres_madre")
            estudiante.dni_madre = form.get("dni_madre")
            estudiante.grado_instruccion_madre = form.get("grado_instruccion_madre")
            estudiante.fecha_nacimiento_madre = form.get("fecha_nacimiento_madre")
            estudiante.direccion_madre = form.get("direccion_madre")
            estudiante.distrito_madre = form.get("distrito_madre")
            estudiante.provincia_madre = form.get("provincia_madre")
            estudiante.celular_madre = form.get("celular_madre")
            estudiante.telefono_fijo_madre = form.get("telefono_fijo_madre")
            estudiante.ocupacion_madre = form.get("ocupacion_madre")
            estudiante.centro_trabajo_madre = form.get("centro_trabajo_madre")
            estudiante.telefono_trabajo_madre = form.get("telefono_trabajo_madre")
            estudiante.estado_civil_madre = form.get("estado_civil_madre")
            estudiante.religion_madre = form.get("religion_madre")
            estudiante.vive_con_hijo_madre = form.get("vive_con_hijo_madre")
            estudiante.correo_madre = form.get("correo_madre")

            estudiante.padre_supervivencia = form.get("padre_supervivencia")
            estudiante.madre_supervivencia = form.get("madre_supervivencia")

            estudiante.apellido_paterno_apoderado = form.get("apellido_paterno_apoderado")
            estudiante.apellido_materno_apoderado = form.get("apellido_materno_apoderado")
            estudiante.nombres_apoderado = form.get("nombres_apoderado")
            estudiante.dni_apoderado = form.get("dni_apoderado")
            estudiante.relacion_apoderado = form.get("relacion_apoderado")
            estudiante.grado_instruccion_apoderado = form.get("grado_instruccion_apoderado")
            estudiante.fecha_nacimiento_apoderado = form.get("fecha_nacimiento_apoderado")
            estudiante.direccion_apoderado = form.get("direccion_apoderado")
            estudiante.distrito_apoderado = form.get("distrito_apoderado")
            estudiante.provincia_apoderado = form.get("provincia_apoderado")
            estudiante.celular_apoderado = form.get("celular_apoderado")
            estudiante.telefono_fijo_apoderado = form.get("telefono_fijo_apoderado")
            estudiante.ocupacion_apoderado = form.get("ocupacion_apoderado")
            estudiante.centro_trabajo_apoderado = form.get("centro_trabajo_apoderado")
            estudiante.telefono_trabajo_apoderado = form.get("telefono_trabajo_apoderado")
            estudiante.estado_civil_apoderado = form.get("estado_civil_apoderado")
            estudiante.religion_apoderado = form.get("religion_apoderado")
            estudiante.vive_con_hijo_apoderado = form.get("vive_con_hijo_apoderado")
            estudiante.correo_apoderado = form.get("correo_apoderado")

            foto = form.get("foto_perfil")
            if isinstance(foto, UploadFile) and foto.filename:
                foto_path = await save_student_photo_upload(foto, estudiante.dni_est)
                if foto_path:
                    estudiante.foto_perfil = foto_path

            db.session.commit()
            add_flash(request, "Estudiante actualizado exitosamente", "success")
            return RedirectResponse(url=str(request.url_for("estudiantes.ver", id=estudiante.id)), status_code=303)

        except Exception as e:
            db.session.rollback()
            print(f"Error al actualizar estudiante: {e}")
            add_flash(request, "Error al actualizar estudiante. Intente nuevamente.", "error")

    return templates.TemplateResponse(
        "editar_estudiante.html",
        common_context(
            request,
            estudiante=estudiante,
            grados_db=_grados_activos_por_nivel(),
        ),
    )


@router.post("/estudiante/{id}/eliminar", name="estudiantes.eliminar")
async def eliminar(
    request: Request, id: int, _user_id: int = Depends(get_current_user_id)
):
    """Eliminar estudiante"""
    form = await request.form()
    if getattr(Config, "WTF_CSRF_ENABLED", True) and not csrf_ok(request, form.get("csrf_token")):
        add_flash(request, "Sesión de seguridad expirada. Intente de nuevo.", "error")
        return RedirectResponse(url=str(request.url_for("estudiantes.list")), status_code=303)
    try:
        estudiante = _get_estudiante_or_404(id)
        nombre_completo = f"{estudiante.nombres_est} {estudiante.apellido_paterno_est}"

        db.session.delete(estudiante)
        db.session.commit()

        add_flash(request, f"Estudiante {nombre_completo} eliminado exitosamente", "success")
    except HTTPException:
        raise
    except Exception as e:
        db.session.rollback()
        print(f"Error al eliminar estudiante: {e}")
        add_flash(request, "Error al eliminar estudiante. Intente nuevamente.", "error")

    return RedirectResponse(url=str(request.url_for("estudiantes.list")), status_code=303)


@router.post("/estudiante/{id}/suspender-acceso", name="estudiantes.toggle_acceso")
async def toggle_acceso(
    request: Request, id: int, _user_id: int = Depends(get_current_user_id)
):
    """Suspende o reactiva el acceso del estudiante al portal (p. ej. por pensión)."""
    form = await request.form()
    page = (form.get("page") or "").strip()
    q = (form.get("q") or "").strip()
    filtro = (form.get("filtro") or "").strip()
    list_url = str(request.url_for("estudiantes.list"))
    params = {k: v for k, v in (("page", page), ("q", q), ("filtro", filtro)) if v}
    if params:
        list_url = f"{list_url}?{urlencode(params)}"

    if getattr(Config, "WTF_CSRF_ENABLED", True) and not csrf_ok(request, form.get("csrf_token")):
        add_flash(request, "Sesión de seguridad expirada. Intente de nuevo.", "error")
        return RedirectResponse(url=list_url, status_code=303)
    try:
        estudiante = _get_estudiante_or_404(id)
        estudiante.acceso_suspendido = not bool(estudiante.acceso_suspendido)
        db.session.commit()
        nombre = f"{estudiante.nombres_est} {estudiante.apellido_paterno_est}"
        if estudiante.acceso_suspendido:
            add_flash(request, f"Acceso al portal suspendido para {nombre}.", "success")
        else:
            add_flash(request, f"Acceso al portal reactivado para {nombre}.", "success")
    except HTTPException:
        raise
    except Exception as e:
        db.session.rollback()
        print(f"Error al cambiar el acceso del estudiante: {e}")
        add_flash(request, "Error al cambiar el acceso del estudiante. Intente nuevamente.", "error")

    return RedirectResponse(url=list_url, status_code=303)


@router.post("/estudiantes/acceso-lote", name="estudiantes.acceso_lote")
async def acceso_lote(request: Request, _user_id: int = Depends(get_current_user_id)):
    """Suspende o reactiva el acceso al portal de varios estudiantes a la vez.

    ``alcance=seleccion`` usa las casillas marcadas; ``alcance=sin_matricula``
    abarca a todos los que tienen matrícula pero ninguna del ciclo en curso,
    sin depender de la paginación.
    """
    form = await request.form()
    page = (form.get("page") or "").strip()
    q = (form.get("q") or "").strip()
    filtro = (form.get("filtro") or "").strip()
    list_url = str(request.url_for("estudiantes.list"))
    params = {k: v for k, v in (("page", page), ("q", q), ("filtro", filtro)) if v}
    if params:
        list_url = f"{list_url}?{urlencode(params)}"

    if getattr(Config, "WTF_CSRF_ENABLED", True) and not csrf_ok(request, form.get("csrf_token")):
        add_flash(request, "Sesión de seguridad expirada. Intente de nuevo.", "error")
        return RedirectResponse(url=list_url, status_code=303)

    accion = (form.get("accion") or "").strip()
    if accion not in ("suspender", "reactivar"):
        add_flash(request, "Acción no válida.", "error")
        return RedirectResponse(url=list_url, status_code=303)
    suspender = accion == "suspender"
    alcance = (form.get("alcance") or "seleccion").strip()

    try:
        anio_ciclo = anio_academico_actual()
        if alcance == "sin_matricula":
            query = _query_sin_matricula_vigente(Estudiante.query, anio_ciclo)
            detalle = f"sin matrícula vigente del ciclo {anio_ciclo}"
        else:
            ids = []
            for raw in form.getlist("estudiante_ids"):
                try:
                    ids.append(int(raw))
                except (TypeError, ValueError):
                    continue
            if not ids:
                add_flash(request, "No seleccionaste ningún estudiante.", "warning")
                return RedirectResponse(url=list_url, status_code=303)
            query = Estudiante.query.filter(Estudiante.id.in_(ids))
            detalle = "seleccionado(s)"

        cambiados = 0
        for estudiante in query.all():
            if bool(estudiante.acceso_suspendido) != suspender:
                estudiante.acceso_suspendido = suspender
                cambiados += 1
        db.session.commit()

        if cambiados:
            verbo = "suspendido" if suspender else "reactivado"
            add_flash(
                request,
                f"Acceso al portal {verbo} para {cambiados} estudiante(s) {detalle}.",
                "success",
            )
        else:
            add_flash(request, "Ningún estudiante cambió de estado.", "warning")
    except Exception as e:  # noqa: BLE001
        db.session.rollback()
        print(f"Error al cambiar accesos en lote: {e}")
        add_flash(request, "Error al cambiar los accesos. Intente nuevamente.", "error")

    return RedirectResponse(url=list_url, status_code=303)


@router.api_route("/form_estudiante", methods=["GET", "POST"], name="estudiantes.form")
async def form_estudiante(request: Request, _user_id: int = Depends(get_current_user_id)):
    """Formulario de registro de estudiante"""
    if request.method == "POST":
        form = await request.form()
        if getattr(Config, "WTF_CSRF_ENABLED", True) and not csrf_ok(request, form.get("csrf_token")):
            add_flash(request, "Sesión de seguridad expirada. Intente de nuevo.", "error")
            return RedirectResponse(url=str(request.url_for("estudiantes.form")), status_code=303)
        try:
            # Create new student instance - COPIAR TODO EL CÓDIGO DE app.py líneas 379-469
            nuevo_estudiante = Estudiante(
                # Información del Estudiante
                nivel=form.get('nivel'),
                grado=form.get('grado'),
                seccion=form.get('seccion'),
                turno=form.get('turno'),
                carrera_postula=form.get('carrera_postula'),
                area_postula=form.get('area_postula'),
                apellido_paterno_est=form.get('apellido_paterno_est'),
                apellido_materno_est=form.get('apellido_materno_est'),
                nombres_est=form.get('nombres_est'),
                dni_est=form.get('dni_est'),
                correo_est=form.get('correo_est'),
                fecha_nacimiento_est=form.get('fecha_nacimiento_est'),
                direccion_est=form.get('direccion_est'),
                distrito_est=form.get('distrito_est'),
                provincia_est=form.get('provincia_est'),
                referencia_est=form.get('referencia_est'),
                numero_celular_est=form.get('numero_celular_est'),
                telefono_fijo_est=form.get('telefono_fijo_est'),
                religion_est=form.get('religion_est'),
                tiene_hermanos=form.get('tiene_hermanos'),
                numero_hermanos=form.get('numero_hermanos'),
                tiene_computadora=form.get('tiene_computadora'),
                grupo_sanguineo=form.get('grupo_sanguineo'),
                es_alergica=form.get('es_alergica'),
                padece_enfermedad=form.get('padece_enfermedad'),
                tiene_discapacidad=form.get('tiene_discapacidad'),
                # Información del Padre
                apellido_paterno_padre=form.get('apellido_paterno_padre'),
                apellido_materno_padre=form.get('apellido_materno_padre'),
                nombres_padre=form.get('nombres_padre'),
                dni_padre=form.get('dni_padre'),
                grado_instruccion_padre=form.get('grado_instruccion_padre'),
                fecha_nacimiento_padre=form.get('fecha_nacimiento_padre'),
                direccion_padre=form.get('direccion_padre'),
                distrito_padre=form.get('distrito_padre'),
                provincia_padre=form.get('provincia_padre'),
                celular_padre=form.get('celular_padre'),
                telefono_fijo_padre=form.get('telefono_fijo_padre'),
                ocupacion_padre=form.get('ocupacion_padre'),
                centro_trabajo_padre=form.get('centro_trabajo_padre'),
                telefono_trabajo_padre=form.get('telefono_trabajo_padre'),
                estado_civil_padre=form.get('estado_civil_padre'),
                religion_padre=form.get('religion_padre'),
                vive_con_hijo_padre=form.get('vive_con_hijo_padre'),
                correo_padre=form.get('correo_padre'),
                # Información de la Madre
                apellido_paterno_madre=form.get('apellido_paterno_madre'),
                apellido_materno_madre=form.get('apellido_materno_madre'),
                nombres_madre=form.get('nombres_madre'),
                dni_madre=form.get('dni_madre'),
                grado_instruccion_madre=form.get('grado_instruccion_madre'),
                fecha_nacimiento_madre=form.get('fecha_nacimiento_madre'),
                direccion_madre=form.get('direccion_madre'),
                distrito_madre=form.get('distrito_madre'),
                provincia_madre=form.get('provincia_madre'),
                celular_madre=form.get('celular_madre'),
                telefono_fijo_madre=form.get('telefono_fijo_madre'),
                ocupacion_madre=form.get('ocupacion_madre'),
                centro_trabajo_madre=form.get('centro_trabajo_madre'),
                telefono_trabajo_madre=form.get('telefono_trabajo_madre'),
                estado_civil_madre=form.get('estado_civil_madre'),
                religion_madre=form.get('religion_madre'),
                vive_con_hijo_madre=form.get('vive_con_hijo_madre'),
                correo_madre=form.get('correo_madre'),
                # Estado de Supervivencia
                padre_supervivencia=form.get('padre_supervivencia'),
                madre_supervivencia=form.get('madre_supervivencia'),
                # Información del Apoderado
                apellido_paterno_apoderado=form.get('apellido_paterno_apoderado'),
                apellido_materno_apoderado=form.get('apellido_materno_apoderado'),
                nombres_apoderado=form.get('nombres_apoderado'),
                dni_apoderado=form.get('dni_apoderado'),
                relacion_apoderado=form.get('relacion_apoderado'),
                grado_instruccion_apoderado=form.get('grado_instruccion_apoderado'),
                fecha_nacimiento_apoderado=form.get('fecha_nacimiento_apoderado'),
                direccion_apoderado=form.get('direccion_apoderado'),
                distrito_apoderado=form.get('distrito_apoderado'),
                provincia_apoderado=form.get('provincia_apoderado'),
                celular_apoderado=form.get('celular_apoderado'),
                telefono_fijo_apoderado=form.get('telefono_fijo_apoderado'),
                ocupacion_apoderado=form.get('ocupacion_apoderado'),
                centro_trabajo_apoderado=form.get('centro_trabajo_apoderado'),
                telefono_trabajo_apoderado=form.get('telefono_trabajo_apoderado'),
                estado_civil_apoderado=form.get('estado_civil_apoderado'),
                religion_apoderado=form.get('religion_apoderado'),
                vive_con_hijo_apoderado=form.get('vive_con_hijo_apoderado'),
                correo_apoderado=form.get('correo_apoderado')
            )

            # Manejar foto del estudiante
            foto_pf = form.get("foto_perfil")
            if isinstance(foto_pf, UploadFile) and foto_pf.filename:
                dni_est = form.get("dni_est") or "sin_dni"
                foto_path = await save_student_photo_upload(foto_pf, str(dni_est))
                if foto_path:
                    nuevo_estudiante.foto_perfil = foto_path

            nuevo_estudiante.codigo_estudiante = Estudiante.generar_codigo_estudiante()
            db.session.add(nuevo_estudiante)
            db.session.commit()

            add_flash(request, "Estudiante registrado exitosamente", "success")
            return RedirectResponse(url=str(request.url_for("estudiantes.list")), status_code=303)

        except Exception as e:
            db.session.rollback()
            print(f"Error al registrar estudiante: {e}")
            add_flash(request, "Error al registrar estudiante. Intente nuevamente.", "error")

    return templates.TemplateResponse(
        "form_estudiante.html",
        common_context(request, grados_db=_grados_activos_por_nivel()),
    )


@router.api_route("/registro_publico", methods=["GET", "POST"], name="estudiantes.registro_publico")
async def registro_publico(request: Request):
    """Formulario público de registro (sin login requerido)"""
    if request.method == "POST":
        form = await request.form()
        if getattr(Config, "WTF_CSRF_ENABLED", True) and not csrf_ok(request, form.get("csrf_token")):
            add_flash(request, "Sesión de seguridad expirada. Intente de nuevo.", "error")
            return RedirectResponse(url=str(request.url_for("estudiantes.registro_publico")), status_code=303)
        try:
            # Create new student instance
            nuevo_estudiante = Estudiante(
                # Información del Estudiante
                nivel=form.get('nivel'),
                grado=form.get('grado'),
                seccion=form.get('seccion'),
                turno=form.get('turno'),
                carrera_postula=form.get('carrera_postula'),
                area_postula=form.get('area_postula'),
                apellido_paterno_est=form.get('apellido_paterno_est'),
                apellido_materno_est=form.get('apellido_materno_est'),
                nombres_est=form.get('nombres_est'),
                dni_est=form.get('dni_est'),
                correo_est=form.get('correo_est'),
                fecha_nacimiento_est=form.get('fecha_nacimiento_est'),
                direccion_est=form.get('direccion_est'),
                distrito_est=form.get('distrito_est'),
                provincia_est=form.get('provincia_est'),
                referencia_est=form.get('referencia_est'),
                numero_celular_est=form.get('numero_celular_est'),
                telefono_fijo_est=form.get('telefono_fijo_est'),
                religion_est=form.get('religion_est'),
                tiene_hermanos=form.get('tiene_hermanos'),
                numero_hermanos=form.get('numero_hermanos'),
                tiene_computadora=form.get('tiene_computadora'),
                grupo_sanguineo=form.get('grupo_sanguineo'),
                es_alergica=form.get('es_alergica'),
                padece_enfermedad=form.get('padece_enfermedad'),
                tiene_discapacidad=form.get('tiene_discapacidad'),

                # Información del Padre
                apellido_paterno_padre=form.get('apellido_paterno_padre'),
                apellido_materno_padre=form.get('apellido_materno_padre'),
                nombres_padre=form.get('nombres_padre'),
                dni_padre=form.get('dni_padre'),
                grado_instruccion_padre=form.get('grado_instruccion_padre'),
                fecha_nacimiento_padre=form.get('fecha_nacimiento_padre'),
                direccion_padre=form.get('direccion_padre'),
                distrito_padre=form.get('distrito_padre'),
                provincia_padre=form.get('provincia_padre'),
                celular_padre=form.get('celular_padre'),
                telefono_fijo_padre=form.get('telefono_fijo_padre'),
                ocupacion_padre=form.get('ocupacion_padre'),
                centro_trabajo_padre=form.get('centro_trabajo_padre'),
                telefono_trabajo_padre=form.get('telefono_trabajo_padre'),
                estado_civil_padre=form.get('estado_civil_padre'),
                religion_padre=form.get('religion_padre'),
                vive_con_hijo_padre=form.get('vive_con_hijo_padre'),
                correo_padre=form.get('correo_padre'),

                # Información de la Madre
                apellido_paterno_madre=form.get('apellido_paterno_madre'),
                apellido_materno_madre=form.get('apellido_materno_madre'),
                nombres_madre=form.get('nombres_madre'),
                dni_madre=form.get('dni_madre'),
                grado_instruccion_madre=form.get('grado_instruccion_madre'),
                fecha_nacimiento_madre=form.get('fecha_nacimiento_madre'),
                direccion_madre=form.get('direccion_madre'),
                distrito_madre=form.get('distrito_madre'),
                provincia_madre=form.get('provincia_madre'),
                celular_madre=form.get('celular_madre'),
                telefono_fijo_madre=form.get('telefono_fijo_madre'),
                ocupacion_madre=form.get('ocupacion_madre'),
                centro_trabajo_madre=form.get('centro_trabajo_madre'),
                telefono_trabajo_madre=form.get('telefono_trabajo_madre'),
                estado_civil_madre=form.get('estado_civil_madre'),
                religion_madre=form.get('religion_madre'),
                vive_con_hijo_madre=form.get('vive_con_hijo_madre'),
                correo_madre=form.get('correo_madre'),

                # Estado de Supervivencia
                padre_supervivencia=form.get('padre_supervivencia'),
                madre_supervivencia=form.get('madre_supervivencia'),

                # Información del Apoderado
                apellido_paterno_apoderado=form.get('apellido_paterno_apoderado'),
                apellido_materno_apoderado=form.get('apellido_materno_apoderado'),
                nombres_apoderado=form.get('nombres_apoderado'),
                dni_apoderado=form.get('dni_apoderado'),
                relacion_apoderado=form.get('relacion_apoderado'),
                grado_instruccion_apoderado=form.get('grado_instruccion_apoderado'),
                fecha_nacimiento_apoderado=form.get('fecha_nacimiento_apoderado'),
                direccion_apoderado=form.get('direccion_apoderado'),
                distrito_apoderado=form.get('distrito_apoderado'),
                provincia_apoderado=form.get('provincia_apoderado'),
                celular_apoderado=form.get('celular_apoderado'),
                telefono_fijo_apoderado=form.get('telefono_fijo_apoderado'),
                ocupacion_apoderado=form.get('ocupacion_apoderado'),
                centro_trabajo_apoderado=form.get('centro_trabajo_apoderado'),
                telefono_trabajo_apoderado=form.get('telefono_trabajo_apoderado'),
                estado_civil_apoderado=form.get('estado_civil_apoderado'),
                religion_apoderado=form.get('religion_apoderado'),
                vive_con_hijo_apoderado=form.get('vive_con_hijo_apoderado'),
                correo_apoderado=form.get('correo_apoderado')
            )

            # Manejar foto del estudiante
            foto_pf = form.get("foto_perfil")
            if isinstance(foto_pf, UploadFile) and foto_pf.filename:
                dni_est = form.get("dni_est") or "sin_dni"
                foto_path = await save_student_photo_upload(foto_pf, str(dni_est))
                if foto_path:
                    nuevo_estudiante.foto_perfil = foto_path

            nuevo_estudiante.codigo_estudiante = Estudiante.generar_codigo_estudiante()
            db.session.add(nuevo_estudiante)
            db.session.commit()

            add_flash(request, "Registro completado exitosamente", "success")
            return RedirectResponse(url=str(request.url_for("estudiantes.registro_exitoso")), status_code=303)

        except Exception as e:
            db.session.rollback()
            print(f"Error al registrar: {e}")
            add_flash(request, "Error al procesar el registro. Por favor, intente nuevamente.", "error")

    return templates.TemplateResponse(
        "formulario.html",
        common_context(request, grados_db=_grados_activos_por_nivel()),
    )


@router.get("/registro_exitoso", name="estudiantes.registro_exitoso")
def registro_exitoso(request: Request):
    """Página de confirmación de registro"""
    return templates.TemplateResponse("registro_exitoso.html", common_context(request))


# ================== RUTA DE ASISTENCIA ==================

@router.get("/estudiante/{id}/asistencia", name="estudiantes.asistencia")
def asistencia(
    request: Request, id: int, _user_id: int = Depends(get_current_user_id)
):
    """Ver historial de asistencia de un estudiante"""
    estudiante = _get_estudiante_or_404(id)
    qp = request.query_params

    mes = _qp_int(qp, "mes")
    anio = _qp_int(qp, "anio") or datetime.now().year

    query = Asistencia.query.filter_by(estudiante_id=id)

    if anio:
        query = query.filter(extract("year", Asistencia.fecha_hora) == anio)
    if mes:
        query = query.filter(extract("month", Asistencia.fecha_hora) == mes)

    total_entradas = query.filter(Asistencia.tipo == "ENTRADA").count()
    total_salidas = query.filter(Asistencia.tipo == "SALIDA").count()

    page = _qp_int(qp, "page") or 1
    per_page = 20
    asistencias = paginate_query(
        query.order_by(Asistencia.fecha_hora.desc()),
        page=page,
        per_page=per_page,
        error_out=False,
    )

    anios_disponibles = db.session.query(
        extract("year", Asistencia.fecha_hora).label("anio")
    ).filter_by(estudiante_id=id).distinct().order_by(
        extract("year", Asistencia.fecha_hora).desc()
    ).all()
    anios_disponibles = (
        [int(a.anio) for a in anios_disponibles] if anios_disponibles else [datetime.now().year]
    )

    return templates.TemplateResponse(
        "asistencia_estudiante.html",
        common_context(
            request,
            estudiante=estudiante,
            asistencias=asistencias,
            total_entradas=total_entradas,
            total_salidas=total_salidas,
            mes_filtro=mes,
            anio_filtro=anio,
            anios_disponibles=anios_disponibles,
        ),
    )


@router.get("/estudiante/{id}/asistencia/pdf", name="estudiantes.asistencia_pdf")
def asistencia_pdf(request: Request, id: int, _user_id: int = Depends(get_current_user_id)):
    """Exportar asistencia de un estudiante como PDF con header/footer institucional"""
    import fpdf

    estudiante = _get_estudiante_or_404(id)
    qp = request.query_params
    mes = _qp_int(qp, "mes")
    anio = _qp_int(qp, "anio") or datetime.now().year

    query = Asistencia.query.filter_by(estudiante_id=id)
    if anio:
        query = query.filter(extract("year", Asistencia.fecha_hora) == anio)
    if mes:
        query = query.filter(extract("month", Asistencia.fecha_hora) == mes)

    asistencias = query.order_by(Asistencia.fecha_hora.asc()).all()

    total_entradas = sum(1 for a in asistencias if a.tipo == 'ENTRADA')
    total_salidas = sum(1 for a in asistencias if a.tipo == 'SALIDA')

    meses_nombres = {
        1: 'Enero', 2: 'Febrero', 3: 'Marzo', 4: 'Abril',
        5: 'Mayo', 6: 'Junio', 7: 'Julio', 8: 'Agosto',
        9: 'Septiembre', 10: 'Octubre', 11: 'Noviembre', 12: 'Diciembre'
    }
    periodo = f"{meses_nombres[mes]} {anio}" if mes else str(anio)

    # --- Constantes de diseño (mismo patrón que academia) ---
    HEADER_HEIGHT_MM = 30
    FOOTER_HEIGHT_MM = 25
    MARGIN_LR = 10
    MARGIN_TOP = HEADER_HEIGHT_MM + 15
    MARGIN_BOTTOM = FOOTER_HEIGHT_MM + 5
    PURPLE_RGB = (69, 3, 140)
    WHITE_RGB = (255, 255, 255)
    LIGHT_GREY_RGB = (240, 240, 240)
    DARK_GREY_RGB = (50, 50, 50)

    # Rutas de imágenes
    header_img = os.path.join(str(ROOT_DIR), 'academia', 'static', 'img', 'encabezado.png')
    footer_img = os.path.join(str(ROOT_DIR), 'academia', 'static', 'img', 'pie.png')
    has_header = os.path.exists(header_img)
    has_footer = os.path.exists(footer_img)

    def _encode(text):
        try:
            return str(text).encode('latin-1', 'replace').decode('latin-1')
        except Exception:
            return "?"

    # --- Clase PDF con header/footer institucional ---
    class PDFAsistencia(fpdf.FPDF):
        def header(self):
            if has_header:
                try:
                    self.image(header_img, x=0, y=0, w=self.w, h=HEADER_HEIGHT_MM)
                except Exception:
                    pass
            self.set_y(MARGIN_TOP)

        def footer(self):
            if has_footer:
                try:
                    self.image(footer_img, x=0, y=self.h - FOOTER_HEIGHT_MM, w=self.w, h=FOOTER_HEIGHT_MM)
                except Exception:
                    pass
            self.set_y(-15)
            self.set_font('Arial', 'I', 7)
            self.set_text_color(128)
            self.cell(0, 5, _encode(f'Generado el {datetime.now().strftime("%d/%m/%Y %H:%M")}'), 0, 0, 'L')
            self.cell(0, 5, _encode(f'Pagina {self.page_no()}/{{nb}}'), 0, 0, 'R')

    # --- Crear PDF ---
    pdf = PDFAsistencia(orientation='P', unit='mm', format='A4')
    pdf.alias_nb_pages()
    pdf.set_margins(MARGIN_LR, MARGIN_TOP, MARGIN_LR)
    pdf.set_auto_page_break(auto=True, margin=MARGIN_BOTTOM)
    pdf.add_page()

    content_w = pdf.w - 2 * MARGIN_LR  # ancho útil

    # --- Título ---
    pdf.set_font('Arial', 'B', 14)
    pdf.set_text_color(*PURPLE_RGB)
    pdf.cell(0, 7, _encode('REPORTE DE ASISTENCIA'), 0, 1, 'C')
    pdf.set_font('Arial', '', 10)
    pdf.set_text_color(*DARK_GREY_RGB)
    pdf.cell(0, 5, _encode(f'Periodo: {periodo}'), 0, 1, 'C')
    pdf.ln(4)

    # --- Info del estudiante (tabla de datos) ---
    pdf.set_font('Arial', 'B', 9)
    pdf.set_fill_color(*PURPLE_RGB)
    pdf.set_text_color(*WHITE_RGB)
    pdf.cell(content_w, 6, _encode('  DATOS DEL ESTUDIANTE'), 1, 1, 'L', 1)

    pdf.set_text_color(*DARK_GREY_RGB)
    label_w = 35
    value_w = content_w / 2 - label_w
    row_h = 6

    info_rows = [
        ('Estudiante:', estudiante.nombre_completo(), 'DNI:', estudiante.dni_est or 'N/A'),
        ('Nivel:', (estudiante.nivel or 'N/A').capitalize(), 'Grado:', f'{estudiante.grado or "N/A"}'),
    ]

    for label1, val1, label2, val2 in info_rows:
        pdf.set_font('Arial', 'B', 8)
        pdf.set_fill_color(*LIGHT_GREY_RGB)
        pdf.cell(label_w, row_h, _encode(label1), 1, 0, 'L', 1)
        pdf.set_font('Arial', '', 8)
        pdf.cell(value_w, row_h, _encode(val1), 1, 0, 'L')
        pdf.set_font('Arial', 'B', 8)
        pdf.cell(label_w, row_h, _encode(label2), 1, 0, 'L', 1)
        pdf.set_font('Arial', '', 8)
        pdf.cell(value_w, row_h, _encode(val2), 1, 1, 'L')

    pdf.ln(4)

    # --- Estadísticas ---
    stat_w = content_w / 3
    pdf.set_font('Arial', 'B', 10)
    pdf.set_fill_color(*PURPLE_RGB)
    pdf.set_text_color(*WHITE_RGB)
    pdf.cell(stat_w, 8, _encode(f'Entradas: {total_entradas}'), 1, 0, 'C', 1)
    pdf.cell(stat_w, 8, _encode(f'Salidas: {total_salidas}'), 1, 0, 'C', 1)
    pdf.cell(stat_w, 8, _encode(f'Total: {len(asistencias)}'), 1, 1, 'C', 1)
    pdf.ln(4)

    # --- Tabla de registros ---
    col_widths = [12, 22, 45, 25, 35, 51]  # N, Fecha, Dia, Hora, Tipo, Observacion
    headers = ['N', 'Fecha', 'Dia', 'Hora', 'Tipo', 'Observacion']

    # Header de tabla
    pdf.set_font('Arial', 'B', 8)
    pdf.set_fill_color(*PURPLE_RGB)
    pdf.set_text_color(*WHITE_RGB)
    for i, h in enumerate(headers):
        pdf.cell(col_widths[i], 7, _encode(h), 1, 0, 'C', 1)
    pdf.ln()

    dias_semana = {
        0: 'Lunes', 1: 'Martes', 2: 'Miercoles', 3: 'Jueves',
        4: 'Viernes', 5: 'Sabado', 6: 'Domingo'
    }

    # Filas de datos con alternancia de color
    pdf.set_text_color(*DARK_GREY_RGB)
    for i, a in enumerate(asistencias, 1):
        # Alternar color de fondo
        if i % 2 == 0:
            pdf.set_fill_color(*LIGHT_GREY_RGB)
            fill = True
        else:
            pdf.set_fill_color(*WHITE_RGB)
            fill = True

        pdf.set_font('Arial', '', 8)
        pdf.cell(col_widths[0], 6, str(i), 1, 0, 'C', fill)
        pdf.cell(col_widths[1], 6, a.fecha_hora.strftime('%d/%m/%Y'), 1, 0, 'C', fill)
        pdf.cell(col_widths[2], 6, _encode(dias_semana.get(a.fecha_hora.weekday(), '')), 1, 0, 'C', fill)
        pdf.cell(col_widths[3], 6, a.fecha_hora.strftime('%H:%M'), 1, 0, 'C', fill)

        # Tipo con color
        if a.tipo == 'ENTRADA':
            pdf.set_fill_color(220, 240, 220)
            pdf.set_text_color(0, 100, 0)
        else:
            pdf.set_fill_color(250, 220, 220)
            pdf.set_text_color(180, 0, 0)
        pdf.set_font('Arial', 'B', 8)
        pdf.cell(col_widths[4], 6, _encode(a.tipo), 1, 0, 'C', 1)

        # Restaurar colores para observación
        pdf.set_text_color(*DARK_GREY_RGB)
        pdf.set_font('Arial', '', 7)
        if i % 2 == 0:
            pdf.set_fill_color(*LIGHT_GREY_RGB)
        else:
            pdf.set_fill_color(*WHITE_RGB)
        pdf.cell(col_widths[5], 6, '', 1, 1, 'L', fill)

    if not asistencias:
        pdf.set_font('Arial', 'I', 9)
        pdf.set_text_color(128)
        pdf.cell(content_w, 10, _encode('No se encontraron registros de asistencia para el periodo seleccionado.'), 1, 1, 'C')

    # --- Output ---
    pdf_content = pdf.output(dest='S')
    if isinstance(pdf_content, str):
        pdf_content = pdf_content.encode('latin-1')

    nombre_archivo = f"asistencia_{estudiante.dni_est or estudiante.id}_{periodo.replace(' ', '_')}.pdf"
    return Response(
        content=pdf_content,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},
    )


# ================== RUTAS DE CARNETS ESTUDIANTILES ==================

@router.get("/carnets", name="estudiantes.seleccionar_carnets")
def seleccionar_carnets(request: Request, _user_id: int = Depends(get_current_user_id)):
    """Interfaz para seleccionar estudiantes para carnet.

    Los filtros (niveles, grados por nivel y secciones) se obtienen de forma
    dinámica desde la base de datos para reflejar solo los valores realmente
    existentes entre los estudiantes registrados.
    """
    niveles_rows = (
        db.session.query(Estudiante.nivel)
        .filter(Estudiante.nivel.isnot(None), Estudiante.nivel != "")
        .distinct()
        .order_by(Estudiante.nivel.asc())
        .all()
    )
    niveles = [r[0] for r in niveles_rows]

    secciones_rows = (
        db.session.query(Estudiante.seccion)
        .filter(Estudiante.seccion.isnot(None), Estudiante.seccion != "")
        .distinct()
        .order_by(Estudiante.seccion.asc())
        .all()
    )
    secciones = [r[0] for r in secciones_rows]

    pares_rows = (
        db.session.query(Estudiante.nivel, Estudiante.grado)
        .filter(
            Estudiante.nivel.isnot(None), Estudiante.nivel != "",
            Estudiante.grado.isnot(None), Estudiante.grado != "",
        )
        .distinct()
        .all()
    )

    def _orden_grado(g: str):
        g = (g or "").strip()
        return (0, int(g)) if g.isdigit() else (1, g.lower())

    grados_por_nivel: dict[str, list[dict[str, str]]] = {}
    for nv, gr in pares_rows:
        if not nv or not gr:
            continue
        key = nv.strip().lower()
        label = f"{gr}°" if gr.strip().isdigit() else gr
        grados_por_nivel.setdefault(key, []).append({"value": gr, "text": label})
    for key in grados_por_nivel:
        grados_por_nivel[key].sort(key=lambda item: _orden_grado(item["value"]))

    return templates.TemplateResponse(
        "carnets/seleccionar_estudiantes.html",
        common_context(
            request,
            niveles=niveles,
            secciones=secciones,
            grados_por_nivel=grados_por_nivel,
        ),
    )


@router.post("/carnets/generar", name="estudiantes.generar_carnets_pdf")
async def generar_carnets_pdf(
    request: Request, _user_id: int = Depends(get_current_user_id)
):
    """Genera PDF con carnets seleccionados"""
    try:
        form = await request.form()
        if getattr(Config, "WTF_CSRF_ENABLED", True) and not csrf_ok(request, form.get("csrf_token")):
            add_flash(request, "Sesión de seguridad expirada. Intente de nuevo.", "error")
            return RedirectResponse(
                url=str(request.url_for("estudiantes.seleccionar_carnets")), status_code=303
            )

        nivel = form.get("nivel")
        grado = form.get("grado")
        secciones = form.getlist("secciones")
        todos = form.get("todos") == "on"
        busqueda = (form.get("busqueda") or "").strip()

        query = Estudiante.query

        if not todos:
            if nivel and nivel != "todos":
                query = query.filter(func.upper(Estudiante.nivel) == nivel.upper())
            if grado and grado != "todos":
                query = query.filter(func.upper(Estudiante.grado) == grado.upper())
            if secciones:
                secciones_upper = [s.upper() for s in secciones]
                query = query.filter(func.upper(Estudiante.seccion).in_(secciones_upper))

        if busqueda:
            query = query.filter(
                (Estudiante.nombres_est.ilike(f"%{busqueda}%"))
                | (Estudiante.apellido_paterno_est.ilike(f"%{busqueda}%"))
                | (Estudiante.apellido_materno_est.ilike(f"%{busqueda}%"))
                | (Estudiante.dni_est.like(f"%{busqueda}%"))
            )

        estudiantes = query.all()

        if not estudiantes:
            add_flash(request, "No hay estudiantes para generar carnets", "warning")
            return RedirectResponse(
                url=str(request.url_for("estudiantes.seleccionar_carnets")), status_code=303
            )

        logo_path = os.path.join(str(ROOT_DIR), "static/img/nk.png")
        buffer = generar_carnets_a4(estudiantes, logo_path, incluir_reverso=False, root_path=str(ROOT_DIR))

        download_name = f"carnets_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"

        return Response(
            content=buffer.getvalue(),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{download_name}"'},
        )

    except Exception as e:
        print(f"Error generando carnets: {e}")
        add_flash(request, "Error al generar los carnets", "error")
        return RedirectResponse(
            url=str(request.url_for("estudiantes.seleccionar_carnets")), status_code=303
        )


@router.get("/estudiante/{id}/carnet/preview", name="estudiantes.carnet_preview")
def carnet_preview(request: Request, id: int, _user_id: int = Depends(get_current_user_id)):
    """Vista previa HTML del carnet individual"""
    estudiante = _get_estudiante_or_404(id)
    return templates.TemplateResponse(
        "carnets/carnet_preview.html", common_context(request, estudiante=estudiante)
    )


@router.get("/estudiante/{id}/carnet/pdf", name="estudiantes.carnet_pdf")
def carnet_pdf(request: Request, id: int, _user_id: int = Depends(get_current_user_id)):
    """Genera PDF del carnet individual"""
    try:
        estudiante = _get_estudiante_or_404(id)
        logo_path = os.path.join(str(ROOT_DIR), "static/img/nk.png")
        buffer = generar_carnets_a4(
            [estudiante],
            logo_path=logo_path,
            incluir_reverso=False,
            root_path=str(ROOT_DIR),
        )
        nombres_archivo = f"{estudiante.apellido_paterno_est}_{estudiante.nombres_est}".replace(" ", "_")
        download_name = f"carnet_{nombres_archivo}.pdf"
        return Response(
            content=buffer.getvalue(),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{download_name}"'},
        )
    except HTTPException:
        raise
    except Exception as e:
        print(f"Error generando PDF de carnet: {e}")
        add_flash(request, "Error al generar el PDF del carnet", "error")
        return RedirectResponse(
            url=str(request.url_for("estudiantes.carnet_preview", id=id)), status_code=303
        )


@router.get("/estudiante/{id}/carnet/imagen/{formato}", name="estudiantes.carnet_imagen")
def carnet_imagen(
    request: Request, id: int, formato: str, _user_id: int = Depends(get_current_user_id)
):
    """Genera imagen PNG o JPG del carnet de un estudiante"""
    try:
        estudiante = _get_estudiante_or_404(id)
        logo_path = os.path.join(str(ROOT_DIR), "static/img/nk.png")

        formato_upper = formato.upper()
        if formato_upper not in ("PNG", "JPG", "JPEG"):
            add_flash(request, "Formato no válido. Use PNG o JPG", "error")
            return RedirectResponse(
                url=str(request.url_for("estudiantes.carnet_preview", id=id)), status_code=303
            )

        buffer = generar_carnet_imagen(
            estudiante, logo_path, formato_upper, root_path=str(ROOT_DIR)
        )

        mimetype = "image/png" if formato_upper == "PNG" else "image/jpeg"

        nombres_archivo = f"{estudiante.apellido_paterno_est}_{estudiante.nombres_est}".replace(" ", "_")
        download_name = f"carnet_{nombres_archivo}.{formato.lower()}"

        return Response(
            content=buffer.getvalue(),
            media_type=mimetype,
            headers={"Content-Disposition": f'attachment; filename="{download_name}"'},
        )
    except HTTPException:
        raise
    except Exception as e:
        print(f"Error generando imagen: {e}")
        add_flash(request, "Error al generar la imagen del carnet", "error")
        return RedirectResponse(
            url=str(request.url_for("estudiantes.carnet_preview", id=id)), status_code=303
        )


@router.get("/estudiantes/descargar-plantilla", name="estudiantes.descargar_plantilla")
def descargar_plantilla(request: Request, _user_id: int = Depends(get_current_user_id)):
    """Descarga una plantilla Excel para carga masiva de estudiantes"""
    try:
        # Crear un nuevo workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Estudiantes"

        # Definir encabezados
        headers = [
            'Apellido Paterno',
            'Apellido Materno',
            'Nombres',
            'DNI',
            'Nivel',
            'Grado/Programa',
            'Carrera al que Postula',
            'Celular Estudiante',
            'Celular Padre',
            'Correo del Padre',
            'Área a la que Postula'
        ]

        # Estilo para encabezados
        header_fill = PatternFill(start_color="5F2A5D", end_color="5F2A5D", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_alignment = Alignment(horizontal="center", vertical="center")

        # Escribir encabezados
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        # Agregar ejemplos en la segunda fila
        ejemplos = [
            'García',
            'Pérez',
            'Juan Carlos',
            '12345678',
            'secundaria',
            '3 / INTENSIVO',
            'Enfermería',
            '912345678',
            '987654321',
            'padre@email.com',
            'Area 1'
        ]

        for col_num, ejemplo in enumerate(ejemplos, 1):
            cell = ws.cell(row=2, column=col_num)
            cell.value = ejemplo
            cell.alignment = Alignment(horizontal="left", vertical="center")

        # Ajustar ancho de columnas
        column_widths = [18, 18, 25, 12, 15, 12, 25, 20, 20, 30]
        for col_num, width in enumerate(column_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = width

        # Guardar en memoria
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        # Usar Response en lugar de send_file para evitar error de fileno en producción
        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="plantilla_estudiantes.xlsx"'},
        )

    except Exception as e:
        print(f"Error generando plantilla: {e}")
        add_flash(request, "Error al generar la plantilla", "error")
        return RedirectResponse(url=str(request.url_for("estudiantes.list")), status_code=303)


@router.post("/estudiantes/upload-excel", name="estudiantes.upload_excel")
async def upload_excel(request: Request, _user_id: int = Depends(get_current_user_id)):
    """Procesa el archivo Excel y carga estudiantes masivamente"""
    try:
        form = await request.form()
        if getattr(Config, "WTF_CSRF_ENABLED", True) and not csrf_ok(request, form.get("csrf_token")):
            return JSONResponse(
                {"success": False, "message": "Sesión de seguridad expirada. Intente de nuevo."},
                status_code=400,
            )

        print("=== INICIO DE UPLOAD EXCEL ===")

        file = _excel_upload_from_form(form)
        if file is None:
            keys = list(form.keys())
            print(f"ERROR: sin archivo en multipart. Claves recibidas: {keys}")
            return JSONResponse(
                {
                    "success": False,
                    "message": "No se recibió el archivo. Seleccione un .xlsx o .xls e intente de nuevo.",
                },
                status_code=400,
            )

        print(f"Archivo recibido: {file.filename}")

        if not file.filename:
            print("ERROR: El archivo no tiene nombre")
            return JSONResponse(
                {"success": False, "message": "No se seleccionó ningún archivo"},
                status_code=400,
            )

        if not str(file.filename).lower().endswith((".xlsx", ".xls")):
            print(f"ERROR: Extensión no válida: {file.filename}")
            return JSONResponse(
                {
                    "success": False,
                    "message": "Formato de archivo no válido. Use .xlsx o .xls",
                },
                status_code=400,
            )

        print("Intentando leer el archivo Excel...")
        raw = await file.read()
        wb = openpyxl.load_workbook(BytesIO(raw))
        ws = wb.active
        print(f"Excel cargado. Filas: {ws.max_row}, Columnas: {ws.max_column}")

        # Validar que tiene datos
        if ws.max_row < 2:
            print("ERROR: El archivo no tiene datos (menos de 2 filas)")
            return JSONResponse(
                {"success": False, "message": "El archivo está vacío o solo tiene encabezados"},
                status_code=400,
            )

        estudiantes_creados = 0
        errores = []
        creados_refs: list[tuple[int, Estudiante]] = []

        # Cargar todos los DNIs existentes en memoria (1 query en vez de N)
        dnis_existentes = set(
            row[0] for row in
            db.session.query(Estudiante.dni_est).filter(Estudiante.dni_est.isnot(None)).all()
        )

        # Códigos únicos en el lote: generar_codigo_estudiante() solo ve la BD,
        # no los pendientes de commit; sin esto todas las filas repiten el mismo código.
        prefijo_cod, num_codigo = Estudiante.siguiente_codigo_estudiante_inicial()

        print(f"Procesando filas desde la 2 hasta la {ws.max_row}...")
        for row_num in range(2, ws.max_row + 1):
            try:
                apellido_paterno = _clean_cell(ws.cell(row=row_num, column=1).value)
                apellido_materno = _clean_cell(ws.cell(row=row_num, column=2).value)
                nombres = _clean_cell(ws.cell(row=row_num, column=3).value)
                dni = _clean_cell(ws.cell(row=row_num, column=4).value)
                nivel = _clean_cell(ws.cell(row=row_num, column=5).value)
                grado = _clean_cell(ws.cell(row=row_num, column=6).value)
                carrera_postula = _clean_cell(ws.cell(row=row_num, column=7).value)
                celular_estudiante = _clean_cell(ws.cell(row=row_num, column=8).value)
                celular_padre = _clean_cell(ws.cell(row=row_num, column=9).value)
                correo_padre = _clean_cell(ws.cell(row=row_num, column=10).value)
                area_postula = _clean_cell(ws.cell(row=row_num, column=11).value)

                # Validar datos obligatorios
                if not all([apellido_paterno, apellido_materno, nombres, dni]):
                    errores.append(f"Fila {row_num}: Faltan datos obligatorios (Apellidos, Nombres o DNI)")
                    continue

                # Validar DNI en memoria (O(1) en vez de query por fila)
                if dni in dnis_existentes:
                    errores.append(f"Fila {row_num}: El DNI {dni} ya existe en el sistema")
                    continue

                # Marcar DNI como usado para prevenir duplicados dentro del mismo Excel
                dnis_existentes.add(dni)

                nuevo_estudiante = Estudiante(
                    codigo_estudiante=f"{prefijo_cod}{num_codigo:05d}",
                    apellido_paterno_est=apellido_paterno,
                    apellido_materno_est=apellido_materno,
                    nombres_est=nombres,
                    dni_est=dni,
                    nivel=nivel,
                    grado=grado,
                    carrera_postula=carrera_postula,
                    area_postula=area_postula,
                    numero_celular_est=celular_estudiante,
                    celular_padre=celular_padre,
                    correo_padre=correo_padre,
                )

                db.session.add(nuevo_estudiante)
                num_codigo += 1
                estudiantes_creados += 1
                creados_refs.append((row_num, nuevo_estudiante))

            except Exception as e:
                errores.append(f"Fila {row_num}: {str(e)}")
                continue

        # Guardar cambios
        print(f"\n=== RESUMEN ===")
        print(f"Estudiantes creados: {estudiantes_creados}")
        print(f"Errores: {len(errores)}")

        matriculados = 0
        avisos: list[str] = []
        if estudiantes_creados > 0:
            print("Guardando cambios en la base de datos...")
            db.session.commit()
            print("✓ Cambios guardados")

            # Auto-matriculación: por cada estudiante recién creado, buscar un aula
            # activa cuyo (nivel, grado) coincida case-insensitive con el del Excel.
            # Si hay exactamente una, se matricula. Si hay 0 o varias, se reporta y
            # el usuario debe asignarlo manualmente desde /aulas.
            usuario = request.session.get("username") or "SISTEMA_BULK_UPLOAD"
            for row_num, est in creados_refs:
                if not est.nivel or not est.grado:
                    avisos.append(
                        f"Fila {row_num}: estudiante creado sin nivel/grado, no se matriculó."
                    )
                    continue
                aulas_match = (
                    Aula.query.filter(
                        Aula.activo.is_(True),
                        func.upper(Aula.nivel) == est.nivel.strip().upper(),
                        func.upper(Aula.grado) == est.grado.strip().upper(),
                    ).all()
                )
                if not aulas_match:
                    avisos.append(
                        f"Fila {row_num}: no hay aula activa para "
                        f"{est.nivel}/{est.grado}; estudiante creado sin matrícula."
                    )
                    continue
                if len(aulas_match) > 1:
                    avisos.append(
                        f"Fila {row_num}: {len(aulas_match)} aulas coinciden con "
                        f"{est.nivel}/{est.grado}; matricular manualmente."
                    )
                    continue
                aula = aulas_match[0]
                _, err = AulaService.matricular_estudiante(
                    estudiante_id=est.id,
                    aula_id=aula.id,
                    anio_escolar=aula.anio_escolar,
                    observaciones="Matrícula automática (carga masiva Excel)",
                    usuario_registro=usuario,
                )
                if err:
                    avisos.append(f"Fila {row_num}: {err}")
                else:
                    matriculados += 1

        # Preparar mensaje de respuesta
        if estudiantes_creados > 0 and len(errores) == 0:
            mensaje = (
                f"Se importaron exitosamente {estudiantes_creados} estudiante(s); "
                f"{matriculados} matriculado(s) automáticamente."
            )
            return JSONResponse(
                {"success": True, "message": mensaje, "avisos": avisos},
                status_code=200,
            )
        if estudiantes_creados > 0 and len(errores) > 0:
            mensaje = (
                f"Se importaron {estudiantes_creados} estudiante(s) "
                f"({matriculados} matriculado(s)). {len(errores)} fila(s) con errores"
            )
            return JSONResponse(
                {"success": True, "message": mensaje, "errores": errores, "avisos": avisos},
                status_code=200,
            )
        mensaje = f"No se pudo importar ningún estudiante. {len(errores)} error(es) encontrado(s)"
        return JSONResponse(
            {"success": False, "message": mensaje, "errores": errores},
            status_code=400,
        )

    except Exception as e:
        db.session.rollback()
        print(f"Error procesando Excel: {e}")
        return JSONResponse(
            {"success": False, "message": f"Error al procesar el archivo: {str(e)}"},
            status_code=500,
        )
